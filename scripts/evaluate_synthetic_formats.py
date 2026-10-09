"""Measure TXT/DOCX/text-PDF parser consistency on the synthetic corpus.

This round-trip check builds each DOCX and text PDF from the same synthetic
reference used by the TXT path. It checks extraction/profile consistency, not
independent format accuracy or hiring quality. It writes no files and makes no
network requests.
"""

from __future__ import annotations

import hashlib
import io
import json
import re
import sqlite3
import sys
import textwrap
from pathlib import Path

from docx import Document

ROOT = Path(__file__).resolve().parents[1]
DATASET = ROOT / "data" / "synthetic-cv-32"
SAMPLE_ID_RE = re.compile(r"r\d{5}\Z")
sys.path.insert(0, str(ROOT))

import server  # noqa: E402


def safe_dataset_file(dataset_root: Path, sample_id: str, filename: str) -> Path:
    path = (dataset_root / sample_id / filename).resolve()
    if not path.is_relative_to(dataset_root):
        raise ValueError(f"Path sampel sintetis di luar direktori dataset: {sample_id}")
    if not path.is_file():
        raise ValueError(f"File {filename} tidak tersedia untuk sampel sintetis {sample_id}.")
    return path


def make_docx(text: str) -> bytes:
    document = Document()
    for line in text.splitlines():
        if line.strip():
            document.add_paragraph(line)
    stream = io.BytesIO()
    document.save(stream)
    return stream.getvalue()


def wrapped_pdf_lines(text: str) -> list[str]:
    wrapped_lines = []
    for paragraph in text.splitlines():
        wrapped_lines.extend(
            textwrap.wrap(paragraph, width=100, break_long_words=True, break_on_hyphens=False) or [""]
        )
    return wrapped_lines or [""]


def make_text_pdf(text: str, lines_per_page: int = 65) -> bytes:
    if lines_per_page < 1:
        raise ValueError("PDF fixture lines_per_page harus positif.")
    wrapped_lines = wrapped_pdf_lines(text)
    page_lines = [wrapped_lines[index:index + lines_per_page] for index in range(0, len(wrapped_lines), lines_per_page)]
    page_ids = [3 + 2 * index for index in range(len(page_lines))]
    font_id = 3 + 2 * len(page_lines)
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        f"<< /Type /Pages /Kids [{kids}] /Count {len(page_lines)} >>".encode("ascii"),
    ]
    for index, lines in enumerate(page_lines):
        page_id = page_ids[index]
        content_id = page_id + 1
        commands = [b"BT /F1 8 Tf 48 752 Td 10 TL"]
        for line_index, line in enumerate(lines):
            escaped = line.encode("ascii").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
            commands.append(b"(" + escaped + b") Tj")
            if line_index + 1 < len(lines):
                commands.append(b"T*")
        commands.append(b"ET")
        stream = b"\n".join(commands) + b"\n"
        page = (
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 {font_id} 0 R >> >> /Contents {content_id} 0 R >>"
        ).encode("ascii")
        content = (
            b"<< /Length " + str(len(stream)).encode("ascii") + b" >>\nstream\n"
            + stream + b"endstream"
        )
        objects.extend((page, content))
    objects.append(
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica /Encoding /WinAnsiEncoding >>"
    )

    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(pdf)


def check_pdf_page_attribution() -> bool:
    expected = [
        (1, "synthetic page one marker"),
        (2, "synthetic page two marker"),
    ]
    _, extracted = server.extract_document(
        "page-attribution.pdf",
        make_text_pdf("synthetic page one marker\nsynthetic page two marker", lines_per_page=1),
    )
    actual = [(page_number, text.strip()) for page_number, text in extracted if text.strip()]
    return actual == expected


def check_pdf_ascii_punctuation() -> bool:
    """Pin the generated fixture's ASCII apostrophe mapping across PDF readers."""
    expected = ["Candidate's profile can't be re-ordered."]
    _, extracted = server.extract_document(
        "ascii-punctuation.pdf",
        make_text_pdf("Candidate's profile can't be re-ordered."),
    )
    return extracted_lines(extracted) == expected


def check_evidence_excerpt_focus() -> dict[str, bool | int]:
    """Ensure capped excerpts retain the matched terms and page locator."""
    python_criterion = {"id": "crit_synthetic_python", "label": "Python", "type": "required", "weight": 1.0}
    positive_line = "Early project details. " + ("routine work continued. " * 80) + "Built production APIs with Python."
    positive = server.match_criterion(python_criterion, [(4, positive_line)])

    negative_line = ("Earlier role details. " * 80) + "No experience with Python."
    negative = server.match_criterion(python_criterion, [(5, negative_line)])

    partial_criterion = {
        "id": "crit_synthetic_stakeholder", "label": "stakeholder management",
        "type": "required", "weight": 1.0,
    }
    partial_line = "stakeholder " + ("project delivery. " * 80) + "management"
    partial = server.match_criterion(partial_criterion, [(6, partial_line)])

    sensitive_line = "Name: Synthetic Candidate; Python engineer"
    redacted_only = server.match_criterion(python_criterion, [(7, sensitive_line)])
    redacted_score, _ = server.score_candidate([python_criterion], [(7, sensitive_line)])
    sensitive_profile = server.profile_from_text("Name: Python\nAge: 30")
    merged_header_profile = server.profile_from_text("DOB: 20 Jan 1992 Skills Python")
    sensitive_education_profile = server.profile_from_text("Name: Bachelor of Science\nAge: Master of Arts")
    merged_education_profile = server.profile_from_text("DOB: 20 Jan 1992 Education Bachelor of Science")
    lengths = [len(positive["snippet"]), len(negative["snippet"]), len(partial["snippet"])]
    return {
        "cases_checked": len(lengths) + 4,
        "positive_match_visible": positive["result"] == "matched" and "Python" in positive["snippet"],
        "negative_cue_visible": (
            negative["result"] == "needs_verification"
            and "No experience" in negative["snippet"]
            and "Python" in negative["snippet"]
        ),
        "partial_terms_visible": (
            partial["result"] == "partial"
            and "stakeholder" in partial["snippet"]
            and "management" in partial["snippet"]
        ),
        "page_attribution_preserved": (
            positive["page_number"] == 4
            and negative["page_number"] == 5
            and partial["page_number"] == 6
        ),
        "all_snippets_within_limit": all(length <= 700 for length in lengths),
        "redacted_only_evidence_unscored": (
            redacted_only["result"] == "needs_verification"
            and not redacted_only["snippet"]
            and redacted_score == 0.0
            and "python" not in sensitive_profile["skills"]
        ),
        "merged_sensitive_header_section_retained": "python" in merged_header_profile["skills"],
        "sensitive_education_lines_filtered": not sensitive_education_profile["education_levels_mentioned"],
        "merged_sensitive_education_heading_retained": (
            merged_education_profile["education_levels_mentioned"] == ["Bachelor of Science"]
        ),
        "max_excerpt_length": max(lengths, default=0),
    }


def check_candidate_score_ordering() -> dict[str, bool | int]:
    """Check only deterministic ordering of controlled lexical evidence scores."""
    criteria = [
        {"id": "crit_python", "label": "Python", "type": "required", "weight": 2.0},
        {"id": "crit_stakeholder", "label": "stakeholder management", "type": "preferred", "weight": 1.0},
    ]
    fixtures = {
        "cand_full": [(1, "Python with stakeholder management.")],
        "cand_partial": [(1, "Python with stakeholder planning and management.")],
        "cand_needs_verification": [
            (1, "No experience with Python."), (2, "Led stakeholder management.")
        ],
        "cand_unknown_z": [(1, "Experienced with SQL only.")],
        "cand_unknown_b": [(1, "Experienced with SQL only.")],
        "cand_unknown_a": [(1, "Experienced with SQL only.")],
    }
    scores = {
        candidate_id: server.score_candidate(criteria, pages)[0]
        for candidate_id, pages in fixtures.items()
    }
    score_assignments_match = (
        scores["cand_full"] == 100.0
        and scores["cand_partial"] == 83.3
        and scores["cand_needs_verification"] == 33.3
        and all(scores[candidate_id] == 0.0 for candidate_id in (
            "cand_unknown_a", "cand_unknown_b", "cand_unknown_z"
        ))
    )
    rows = (
        ("cand_unknown_z", scores["cand_unknown_z"], "2026-09-30T10:00:00+00:00"),
        ("cand_unknown_b", scores["cand_unknown_b"], "2026-09-30T09:00:00+00:00"),
        ("cand_partial", scores["cand_partial"], "2026-09-30T09:30:00+00:00"),
        ("cand_unknown_a", scores["cand_unknown_a"], "2026-09-30T09:00:00+00:00"),
        ("cand_needs_verification", scores["cand_needs_verification"], "2026-09-30T09:20:00+00:00"),
        ("cand_full", scores["cand_full"], "2026-09-30T09:40:00+00:00"),
    )
    with sqlite3.connect(":memory:") as db:
        db.execute("CREATE TABLE candidates (id TEXT, job_id TEXT, score REAL, created_at TEXT)")
        db.executemany("INSERT INTO candidates VALUES (?, 'job_synthetic', ?, ?)", rows)
        ordered = [row[0] for row in db.execute(
            f"SELECT id FROM candidates WHERE job_id=? ORDER BY {server.CANDIDATE_ORDER_BY_SQL}",
            ("job_synthetic",),
        )]
    expected = [
        "cand_full", "cand_partial", "cand_needs_verification",
        "cand_unknown_a", "cand_unknown_b", "cand_unknown_z",
    ]
    ordered_scores_nonincreasing = all(
        scores[left] >= scores[right] for left, right in zip(ordered, ordered[1:])
    )
    return {
        "candidate_count": len(fixtures),
        "score_assignments_match": score_assignments_match,
        "ordered_scores_nonincreasing": ordered_scores_nonincreasing,
        "timestamp_and_id_ties_deterministic": ordered == expected,
        "all_checks_passed": score_assignments_match and ordered_scores_nonincreasing and ordered == expected,
    }


def check_docx_table_order() -> dict[str, bool]:
    """Probe interleaved paragraphs and columns with a fully synthetic DOCX."""
    document = Document()
    document.add_paragraph("Synthetic opening marker")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Python"
    table.cell(0, 1).text = "FastAPI"
    table.cell(1, 0).text = "PostgreSQL"
    table.cell(1, 1).text = "Docker"
    document.add_paragraph("Synthetic closing marker")
    stream = io.BytesIO()
    document.save(stream)

    _, pages = server.extract_document("synthetic-layout.docx", stream.getvalue())
    actual_lines = extracted_lines(pages)
    expected_lines = [
        "Synthetic opening marker",
        "Python | FastAPI",
        "PostgreSQL | Docker",
        "Synthetic closing marker",
    ]
    expected_profile = profile_signature("\n".join(expected_lines))
    actual_profile = profile_signature("\n".join(actual_lines))
    return {
        "text_order_matched": actual_lines == expected_lines,
        "profile_matched": actual_profile == expected_profile,
    }


def pdf_ascii_projection(text: str) -> str:
    printable = "".join(char if 32 <= ord(char) <= 126 else " " for char in text)
    return " ".join(printable.split())


def extracted_lines(pages: list[tuple[int | None, str]]) -> list[str]:
    return [line.strip() for _, line in pages if line.strip()]


def normalized_reference_lines(text: str) -> list[str]:
    return [line.strip() for line in text.splitlines() if line.strip()]


def profile_signature(text: str) -> tuple[tuple[str, ...], str | None, tuple[str, ...], int | None]:
    profile = server.profile_from_text(text)
    return (
        tuple(sorted(profile["skills"])),
        profile["experience_duration_text"],
        tuple(profile["education_levels_mentioned"]),
        profile["experience_years_mentioned"],
    )


def evaluate() -> dict:
    dataset_root = DATASET.resolve()
    manifest_path = (dataset_root / "manifest.json").resolve()
    if not manifest_path.is_relative_to(dataset_root) or not manifest_path.is_file():
        raise ValueError("Manifest korpus sintetis tidak tersedia di direktori dataset.")
    manifest_bytes = manifest_path.read_bytes()
    manifest = json.loads(manifest_bytes.decode("utf-8"))
    if not isinstance(manifest, list) or len(manifest) != server.SYNTHETIC_DATASET_SIZE:
        raise ValueError("Manifest dataset sintetis kosong atau ukurannya tidak sesuai.")

    sample_ids: list[str] = []
    seen_ids: set[str] = set()
    for entry in manifest:
        sample_id = str(entry.get("sample_id", "")) if isinstance(entry, dict) else ""
        if not SAMPLE_ID_RE.fullmatch(sample_id) or sample_id in seen_ids:
            raise ValueError("ID sampel sintetis tidak valid atau duplikat di manifest.")
        seen_ids.add(sample_id)
        sample_ids.append(sample_id)

    outcomes = {
        "txt": {"text_match": 0, "profile_match": 0, "mismatches": []},
        "docx": {"text_match": 0, "profile_match": 0, "mismatches": []},
        "text_pdf": {"text_match": 0, "profile_match": 0, "mismatches": []},
    }
    dataset_digest = hashlib.sha256()
    dataset_digest.update(b"manifest\0")
    dataset_digest.update(manifest_bytes)
    for sample_id in sample_ids:
        source_path = safe_dataset_file(dataset_root, sample_id, "resume_text.txt")
        source_bytes = source_path.read_bytes()
        dataset_digest.update(sample_id.encode("ascii") + b"\0")
        dataset_digest.update(len(source_bytes).to_bytes(8, "big"))
        dataset_digest.update(source_bytes)
        source_text = source_bytes.decode("utf-8")
        expected_lines = normalized_reference_lines(source_text)
        expected_profile = profile_signature(source_text)

        txt_type, txt_pages = server.extract_document(f"{sample_id}.txt", source_text.encode("utf-8"))
        txt_lines = extracted_lines(txt_pages)
        txt_profile = profile_signature("\n".join(txt_lines))
        if txt_type == "txt" and txt_lines == expected_lines:
            outcomes["txt"]["text_match"] += 1
        else:
            outcomes["txt"]["mismatches"].append(sample_id)
        if txt_profile == expected_profile:
            outcomes["txt"]["profile_match"] += 1

        docx_payload = make_docx(source_text)
        docx_type, docx_pages = server.extract_document(f"{sample_id}.docx", docx_payload)
        docx_lines = extracted_lines(docx_pages)
        docx_profile = profile_signature("\n".join(docx_lines))
        if docx_type == "docx" and docx_lines == expected_lines:
            outcomes["docx"]["text_match"] += 1
        else:
            outcomes["docx"]["mismatches"].append(sample_id)
        if docx_profile == expected_profile:
            outcomes["docx"]["profile_match"] += 1

        pdf_reference = pdf_ascii_projection(source_text)
        _, pdf_pages = server.extract_document(
            f"{sample_id}.pdf", make_text_pdf(pdf_reference)
        )
        pdf_lines = extracted_lines(pdf_pages)
        pdf_profile = profile_signature("\n".join(pdf_lines))
        if " ".join(" ".join(pdf_lines).split()) == pdf_reference:
            outcomes["text_pdf"]["text_match"] += 1
        else:
            outcomes["text_pdf"]["mismatches"].append(sample_id)
        if pdf_profile == profile_signature(pdf_reference):
            outcomes["text_pdf"]["profile_match"] += 1

    count = len(sample_ids)
    docx_table_probe = check_docx_table_order()
    if not all(docx_table_probe.values()):
        raise ValueError("Probe tata letak tabel DOCX sintetis gagal.")
    pdf_punctuation_probe = check_pdf_ascii_punctuation()
    if not pdf_punctuation_probe:
        raise ValueError("Probe tanda baca ASCII PDF sintetis gagal.")
    evidence_excerpt_probe = check_evidence_excerpt_focus()
    if not all(
        evidence_excerpt_probe[key]
        for key in (
            "positive_match_visible", "negative_cue_visible", "partial_terms_visible",
            "page_attribution_preserved", "all_snippets_within_limit",
            "redacted_only_evidence_unscored", "merged_sensitive_header_section_retained",
            "sensitive_education_lines_filtered", "merged_sensitive_education_heading_retained",
        )
    ):
        raise ValueError("Probe kutipan evidence panjang sintetis gagal.")
    candidate_score_order_probe = check_candidate_score_ordering()
    if not candidate_score_order_probe["all_checks_passed"]:
        raise ValueError("Probe urutan skor evidence kandidat sintetis gagal.")
    evaluator_bytes = Path(__file__).read_bytes().replace(b"\r\n", b"\n")
    return {
        "dataset": "synthetic-cv-32",
        "pipeline_version": server.pipeline_version(),
        "dataset_manifest_sha256": hashlib.sha256(manifest_bytes).hexdigest(),
        "dataset_content_sha256": dataset_digest.hexdigest(),
        "evaluator_sha256": hashlib.sha256(evaluator_bytes).hexdigest(),
        "samples_processed": count,
        "comparison": "TXT/DOCX and generated in-memory single-column text PDF with wrapped lines extracted against the same local synthetic text references; PDF uses a printable-ASCII projection",
        "formats": {
            name: {
                "text_match_count": values["text_match"],
                "text_match_rate": round(values["text_match"] / count, 4) if count else None,
                "profile_match_count": values["profile_match"],
                "profile_match_rate": round(values["profile_match"] / count, 4) if count else None,
                "mismatched_synthetic_sample_ids": values["mismatches"],
            }
            for name, values in outcomes.items()
        },
        "pdf_page_attribution_probe": {
            "expected_pages": 2,
            "matched": check_pdf_page_attribution(),
        },
        "pdf_ascii_punctuation_probe": {"matched": pdf_punctuation_probe},
        "evidence_excerpt_focus_probe": evidence_excerpt_probe,
        "candidate_score_order_probe": candidate_score_order_probe,
        "docx_table_order_probe": docx_table_probe,
        "limitations": [
            "The DOCX and text PDF files are generated directly from the same synthetic text references; this measures parser round-trip consistency, not independent document accuracy.",
            "A separate synthetic DOCX probe checks paragraph/table block order and row-wise two-column flattening; it is not a representative sample of independently authored multi-column CVs.",
            "The generated text PDF is a simple single-column fixture with wrapped lines, not an independently authored CV layout.",
            "The text PDF uses a printable-ASCII projection; non-ASCII/control characters are replaced with spaces and whitespace is collapsed.",
            "It does not evaluate scanned-PDF OCR, visual layout, languages outside the English synthetic corpus, or hiring quality/fairness.",
            "The evidence-excerpt probe uses a controlled synthetic fixture; it does not estimate citation correctness on the corpus or independently authored CVs.",
            "The candidate-order probe checks lexical evidence score ordering and deterministic tie-breaks only; it is not a relevance, shortlist quality, or hiring benchmark.",
            "No candidate text or profile values are printed or persisted; output contains aggregate counts, synthetic sample IDs, and reproducibility fingerprints.",
        ],
    }


def main() -> int:
    try:
        result = evaluate()
    except (OSError, ValueError, json.JSONDecodeError) as error:
        print(f"Evaluasi format sintetis gagal: {error}", file=sys.stderr)
        return 2
    print(json.dumps(result, ensure_ascii=False, indent=2))
    if any(
        values["text_match_count"] != result["samples_processed"]
        or values["profile_match_count"] != result["samples_processed"]
        for values in result["formats"].values()
    ):
        return 1
    if not result["pdf_page_attribution_probe"]["matched"]:
        return 1
    if not all(result["docx_table_order_probe"].values()):
        return 1
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
