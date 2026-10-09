"""Exercise the local API lifecycle using an isolated temporary SQLite database."""

from __future__ import annotations

import io
import http.client
import json
import os
import re
import sqlite3
import socket
import signal
import subprocess
import sys
import tempfile
import threading
import time
import urllib.error
import urllib.request
import zipfile
import csv
from http.server import BaseHTTPRequestHandler
from contextlib import redirect_stdout
from pathlib import Path
from urllib.parse import urlsplit
from unittest.mock import patch

from PIL import Image


ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / "scripts"))

import server  # noqa: E402
import evaluate_synthetic  # noqa: E402
import evaluate_synthetic_ocr as synthetic_ocr  # noqa: E402
import evaluate_synthetic_ocr_corpus as synthetic_ocr_corpus  # noqa: E402
import evaluate_synthetic_formats as synthetic_formats  # noqa: E402
import evaluate_adjudications  # noqa: E402
import evaluate_rankings  # noqa: E402
import annotation_packets  # noqa: E402
import ranking_packets  # noqa: E402
import backup_database  # noqa: E402
import restore_database  # noqa: E402
import preflight_local_release as local_release  # noqa: E402


def request_json(url: str, method: str = "GET", payload=None, headers=None):
    data = json.dumps(payload).encode("utf-8") if payload is not None else None
    request_headers = dict(headers or {})
    if payload is not None:
        request_headers["Content-Type"] = "application/json"
    if method.upper() in {"POST", "PUT", "PATCH", "DELETE"}:
        parsed = urlsplit(url)
        request_headers["Origin"] = f"{parsed.scheme}://{parsed.netloc}"
        request_headers["X-KarsaHire-Request"] = "same-origin-ui"
    request = urllib.request.Request(url, data=data, headers=request_headers, method=method)
    with urllib.request.urlopen(request, timeout=10) as response:
        body = response.read()
        return response.status, json.loads(body.decode("utf-8")) if body else None


def upload_text(url: str, filename: str, text: str):
    return upload_bytes(url, filename, text.encode("utf-8"))


def upload_bytes(url: str, filename: str, payload: bytes):
    boundary = "----KarsaHireSmokeBoundary"
    body = (
        f"--{boundary}\r\n"
        f'Content-Disposition: form-data; name="file"; filename="{filename}"\r\n'
        "Content-Type: application/octet-stream\r\n\r\n"
    ).encode("utf-8") + payload + f"\r\n--{boundary}--\r\n".encode("ascii")
    parsed = urlsplit(url)
    headers = {
        "Content-Type": f"multipart/form-data; boundary={boundary}",
        "Origin": f"{parsed.scheme}://{parsed.netloc}",
        "X-KarsaHire-Request": "same-origin-ui",
    }
    return _request_raw(url, body, headers)


def _request_raw(url: str, body: bytes, headers: dict):
    request = urllib.request.Request(url, data=body, headers=headers, method="POST")
    with urllib.request.urlopen(request, timeout=10) as response:
        return response.status, json.loads(response.read().decode("utf-8"))


def check_request_guards(base: str) -> None:
    with urllib.request.urlopen(base, timeout=10) as response:
        html = response.read().decode("utf-8")
        request_id = response.headers.get("X-Request-ID", "")
        assert re.fullmatch(r"[0-9a-f]{32}", request_id), "response should include a server-generated request ID"
        assert response.headers.get("Server") == "KarsaHire", "server banner should not disclose the Python runtime version"
        assert response.headers.get("X-Content-Type-Options") == "nosniff"
        assert "script-src 'self'" in response.headers.get("Content-Security-Policy", "")
        assert "/app.js" in html

    with urllib.request.urlopen(f"{base}/api/health", timeout=10) as response:
        assert response.headers.get("X-Content-Type-Options") == "nosniff"
        assert response.headers.get("X-Frame-Options") == "DENY"
        assert "frame-ancestors 'none'" in response.headers.get("Content-Security-Policy", "")
        health = json.loads(response.read().decode("utf-8"))
        assert health["pipeline_version"] == server.pipeline_version()
        assert health["ready"] is True and health["schema_version"] == server.SCHEMA_VERSION

    with urllib.request.urlopen(f"{base}/api/health/live", timeout=10) as response:
        assert response.status == 200
        request_id = response.headers.get("X-Request-ID", "")
        assert re.fullmatch(r"[0-9a-f]{32}", request_id)
        live = json.loads(response.read().decode("utf-8"))
        assert live["ok"] is True and live["live"] is True and live["draining"] is False


def check_request_id_correlation(base: str) -> None:
    supplied = urllib.request.Request(
        f"{base}/api/health/live",
        headers={"X-Request-ID": "client-controlled"},
    )
    output = io.StringIO()
    with redirect_stdout(output):
        with urllib.request.urlopen(supplied, timeout=10) as response:
            request_id = response.headers.get("X-Request-ID", "")
            response.read()
    assert re.fullmatch(r"[0-9a-f]{32}", request_id)
    assert request_id != "client-controlled"
    records = [json.loads(line) for line in output.getvalue().splitlines() if line.startswith("{")]
    matching = [
        record for record in records
        if record.get("event") == "http_request" and record.get("route") == "/api/health/live"
    ]
    assert len(matching) == 1 and matching[0].get("request_id") == request_id
    print("PASS: response request ID matches its structured log and ignores a client-supplied value")

    with server.connect() as db:
        db.execute(f"PRAGMA user_version = {server.SCHEMA_VERSION + 1}")
    with urllib.request.urlopen(f"{base}/api/health/live", timeout=10) as response:
        assert response.status == 200, "liveness must not depend on database readiness"
    try:
        urllib.request.urlopen(f"{base}/api/health/ready", timeout=10)
    except urllib.error.HTTPError as error:
        readiness = json.loads(error.read().decode("utf-8"))
        assert error.code == 503 and readiness["ready"] is False
    else:
        raise AssertionError("readiness should fail when the database schema is unsupported")
    with server.connect() as db:
        db.execute(f"PRAGMA user_version = {server.SCHEMA_VERSION}")

    for content_type, body in (("text/plain", None), ("application/json", b"[]")):
        request = urllib.request.Request(
            f"{base}/api/jobs",
            data=body,
            headers={
                "Content-Type": content_type,
                "Origin": base,
                "X-KarsaHire-Request": "same-origin-ui",
            },
            method="POST",
        )
        try:
            urllib.request.urlopen(request, timeout=10)
        except urllib.error.HTTPError as error:
            error.read()
            assert error.code == 400, f"invalid JSON request returned HTTP {error.code}"
            assert error.headers.get("X-Content-Type-Options") == "nosniff"
        else:
            raise AssertionError(f"request with Content-Type {content_type} should be rejected")

    parsed = urlsplit(base)
    framing_cases = (
        (("Content-Length", "0"), ("Content-Length", "0")),
        (("Content-Length", "invalid"),),
        (("Content-Length", "9" * 21),),
        (("Transfer-Encoding", "chunked"),),
    )
    for framing_headers in framing_cases:
        connection = http.client.HTTPConnection(parsed.hostname, parsed.port, timeout=5)
        try:
            connection.putrequest("POST", "/api/jobs", skip_host=True)
            connection.putheader("Host", f"{parsed.hostname}:{parsed.port}")
            connection.putheader("Origin", base)
            connection.putheader("X-KarsaHire-Request", "same-origin-ui")
            connection.putheader("Content-Type", "application/json")
            for name, value in framing_headers:
                connection.putheader(name, value)
            connection.endheaders()
            response = connection.getresponse()
            response.read()
            assert response.status == 400, "unsupported or ambiguous request framing should be rejected"
        finally:
            connection.close()

    duplicate_security_headers = (
        (("Host", f"{parsed.hostname}:{parsed.port}"), ("Host", f"attacker.invalid:{parsed.port}")),
        (("Origin", base), ("Origin", "http://attacker.invalid")),
        (("X-KarsaHire-Request", "same-origin-ui"), ("X-KarsaHire-Request", "forged")),
    )
    for duplicate_headers in duplicate_security_headers:
        connection = http.client.HTTPConnection(parsed.hostname, parsed.port, timeout=5)
        try:
            connection.putrequest("POST", "/api/jobs", skip_host=True)
            connection.putheader("Host", f"{parsed.hostname}:{parsed.port}")
            connection.putheader("Origin", base)
            connection.putheader("X-KarsaHire-Request", "same-origin-ui")
            connection.putheader("Content-Type", "application/json")
            connection.putheader("Content-Length", "0")
            for name, value in duplicate_headers:
                connection.putheader(name, value)
            connection.endheaders()
            response = connection.getresponse()
            response.read()
            assert response.status == 403, "duplicate Host/Origin/UI marker headers should fail closed"
        finally:
            connection.close()


def check_request_body_deadline() -> None:
    class FakeClock:
        def __init__(self):
            self.now = 100.0

        def monotonic(self):
            return self.now

    class FakeConnection:
        def __init__(self):
            self.timeout = 7.0

        def gettimeout(self):
            return self.timeout

        def settimeout(self, value):
            self.timeout = value

    class SlowReader:
        def __init__(self, clock):
            self.clock = clock

        def read(self, _size):
            self.clock.now += 2.0
            return b"x"

    clock = FakeClock()
    handler = type("FakeHandler", (), {})()
    handler.connection = FakeConnection()
    handler.rfile = SlowReader(clock)
    handler.close_connection = False
    with patch.object(server, "time", clock), patch.object(server, "MAX_BODY_READ_SECONDS", 1):
        try:
            server.read_request_body(handler, 2)
        except server.RequestBodyTimeout:
            pass
        else:
            raise AssertionError("request bodies that exceed the total deadline should time out")
    assert handler.close_connection, "timed-out body connection should be closed"
    assert handler.connection.timeout == 7.0, "body reader should restore the socket's original timeout"


def check_incomplete_request_header_timeout(base: str) -> None:
    parsed = urlsplit(base)
    previous_timeout = server.REQUEST_HEADER_IDLE_TIMEOUT_SECONDS
    server.REQUEST_HEADER_IDLE_TIMEOUT_SECONDS = 0.1
    timeout_logs = io.StringIO()
    try:
        with redirect_stdout(timeout_logs):
            with socket.create_connection((parsed.hostname, parsed.port), timeout=2) as connection:
                connection.settimeout(2)
                connection.sendall(
                    f"GET /api/candidates/cand_0123456789ab HTTP/1.1\r\n"
                    f"Host: 127.0.0.1:{parsed.port}\r\n".encode("ascii")
                )
                response = connection.recv(1024)
        assert response == b"", "incomplete request headers should time out and close without a response"
    finally:
        server.REQUEST_HEADER_IDLE_TIMEOUT_SECONDS = previous_timeout

    log = timeout_logs.getvalue()
    assert '"reason":"timeout"' in log
    assert "127.0.0.1" not in log and "cand_0123456789ab" not in log
    assert '"route":"/api/candidates/:id"' in log
    print("PASS: incomplete headers time out and logs redact the peer address and candidate ID")


def check_absolute_request_header_deadline(base: str) -> None:
    parsed = urlsplit(base)
    previous_idle = server.REQUEST_HEADER_IDLE_TIMEOUT_SECONDS
    previous_total = server.REQUEST_HEADER_TOTAL_TIMEOUT_SECONDS
    test_idle_timeout = 1.0
    server.REQUEST_HEADER_IDLE_TIMEOUT_SECONDS = test_idle_timeout
    server.REQUEST_HEADER_TOTAL_TIMEOUT_SECONDS = 0.35
    timeout_logs = io.StringIO()
    pulse_count = 0
    elapsed = 0.0
    try:
        with redirect_stdout(timeout_logs):
            with socket.create_connection((parsed.hostname, parsed.port), timeout=2) as connection:
                connection.settimeout(2)
                start = time.monotonic()
                connection.sendall(
                    f"GET /api/health/live HTTP/1.1\r\nHost: 127.0.0.1:{parsed.port}\r\n".encode("ascii")
                )
                for index in range(12):
                    time.sleep(0.05)
                    try:
                        connection.sendall(f"X-Slow: {index}\r\n".encode("ascii"))
                        pulse_count += 1
                    except OSError:
                        break
                try:
                    response = connection.recv(1024)
                except (ConnectionAbortedError, ConnectionResetError):
                    # Windows may reset a socket when unread partial headers remain on close.
                    response = b""
                elapsed = time.monotonic() - start
    finally:
        server.REQUEST_HEADER_IDLE_TIMEOUT_SECONDS = previous_idle
        server.REQUEST_HEADER_TOTAL_TIMEOUT_SECONDS = previous_total

    log = timeout_logs.getvalue()
    assert response == b"", "absolute header deadline should close without waiting for the idle timeout"
    assert pulse_count >= 3, "client must keep sending before the longer idle timeout"
    assert elapsed < test_idle_timeout, "absolute header deadline should fire before the configured idle timeout"
    assert '"reason":"timeout"' in log
    print("PASS: live loopback slow-header request hits its absolute deadline before idle timeout")


def check_request_header_limits() -> None:
    class StubConnection:
        def __init__(self, payload: bytes, timeout: float = 5.0):
            self.payload = bytearray(payload)
            self.timeout = timeout
            self.recv_calls = 0

        def gettimeout(self):
            return self.timeout

        def settimeout(self, timeout):
            self.timeout = timeout

        def recv(self, size):
            self.recv_calls += 1
            chunk = bytes(self.payload[:size])
            del self.payload[:size]
            return chunk

    deadline_connection = StubConnection(b"")
    deadline_reader = server.RequestHeadReader(deadline_connection)
    with patch.object(server.time, "monotonic", side_effect=(10.0, 26.0)):
        deadline_reader.begin_request()
        try:
            deadline_reader.readline(65_537)
        except socket.timeout:
            pass
        else:
            raise AssertionError("request headers should stop at the absolute deadline")
    assert deadline_connection.recv_calls == 0, "expired header deadline should stop before another socket read"

    oversized_header = (
        b"GET / HTTP/1.1\r\n"
        + b"X-Test: " + b"a" * server.MAX_REQUEST_HEADER_BYTES
        + b"\r\n\r\n"
    )
    limit_reader = server.RequestHeadReader(StubConnection(oversized_header))
    limit_reader.begin_request()
    assert limit_reader.readline(65_537) == b"GET / HTTP/1.1\r\n"
    try:
        limit_reader.readline(65_537)
    except http.client.HTTPException as error:
        assert "aggregate request headers" in str(error)
    else:
        raise AssertionError("aggregate request headers over 64 KiB should be rejected")

    buffered_connection = StubConnection(
        b"POST /api/jobs HTTP/1.1\r\nHost: 127.0.0.1\r\nContent-Length: 4\r\n\r\nDATA"
    )
    buffered_reader = server.RequestHeadReader(buffered_connection)
    buffered_reader.begin_request()
    assert buffered_reader.readline(65_537) == b"POST /api/jobs HTTP/1.1\r\n"
    assert buffered_reader.readline(65_537) == b"Host: 127.0.0.1\r\n"
    assert buffered_reader.readline(65_537) == b"Content-Length: 4\r\n"
    assert buffered_reader.readline(65_537) == b"\r\n"
    assert buffered_reader.read(4) == b"DATA", "buffered body bytes should survive header parsing"
    assert buffered_connection.timeout == 5.0, "header parsing should restore the socket timeout for body reads"
    print("PASS: absolute request-header deadline, aggregate byte cap, and buffered body handoff")


def check_same_origin_mutation_guard(base: str) -> None:
    cases = (
        ({"Origin": base}, "same-origin UI marker should be required"),
        ({"Origin": "http://attacker.invalid", "X-KarsaHire-Request": "same-origin-ui"},
         "cross-origin writes should be rejected"),
        ({"Origin": base, "X-KarsaHire-Request": "same-origin-ui", "Host": "attacker.invalid"},
         "non-loopback Host headers should be rejected"),
    )
    for headers, message in cases:
        request = urllib.request.Request(
            f"{base}/api/jobs",
            data=None,
            headers={"Content-Type": "application/json", **headers},
            method="POST",
        )
        try:
            urllib.request.urlopen(request, timeout=10)
        except urllib.error.HTTPError as error:
            error.read()
            assert error.code == 403, f"{message}; received HTTP {error.code}"
        else:
            raise AssertionError(message)

    request = urllib.request.Request(f"{base}/api/jobs", headers={"Host": "attacker.invalid"})
    try:
        urllib.request.urlopen(request, timeout=10)
    except urllib.error.HTTPError as error:
        error.read()
        assert error.code == 403, f"non-loopback read Host should be rejected; received HTTP {error.code}"
    else:
        raise AssertionError("non-loopback Host headers should be rejected on read endpoints")

    parsed = urlsplit(base)
    connection = http.client.HTTPConnection(parsed.hostname, parsed.port, timeout=2)
    try:
        connection.putrequest("POST", "/api/jobs", skip_host=True)
        connection.putheader("Host", "attacker.invalid")
        connection.putheader("Origin", "http://attacker.invalid")
        connection.putheader("Content-Type", "application/json")
        connection.putheader("Content-Length", "1000000")
        response_finished = threading.Event()
        rejected_connection_closed = []
        send_json = server.Handler.send_json

        def observe_rejection(handler, status, payload, extra_headers=None):
            if status == 403:
                rejected_connection_closed.append(handler.close_connection)
                try:
                    return send_json(handler, status, payload, extra_headers)
                finally:
                    response_finished.set()
            return send_json(handler, status, payload, extra_headers)

        started = time.monotonic()
        with patch.object(server.Handler, "send_json", observe_rejection):
            connection.endheaders()
            assert response_finished.wait(timeout=3), "incomplete hostile mutation should be rejected promptly without draining its body"
        assert time.monotonic() - started < 3, "hostile mutation should close promptly without draining its body"
        assert rejected_connection_closed == [True], "rejected incomplete mutation should close its connection"
    finally:
        connection.close()


def check_local_image_normalization() -> None:
    image_bytes = io.BytesIO()
    Image.new("RGB", (3, 2), color=(42, 91, 66)).save(image_bytes, format="PNG")
    payload = image_bytes.getvalue()
    normalized = server.normalize_image(payload, expected_format="PNG")
    with Image.open(io.BytesIO(normalized)) as image:
        assert image.format == "PNG" and image.size == (3, 2)
    jpeg_bytes = io.BytesIO()
    Image.new("RGB", (3, 2), color=(42, 91, 66)).save(jpeg_bytes, format="JPEG")
    jpeg_normalized = server.normalize_image(jpeg_bytes.getvalue(), expected_format="JPEG")
    with Image.open(io.BytesIO(jpeg_normalized)) as image:
        assert image.format == "PNG" and image.size == (3, 2)
    with patch.object(server, "call_ocr_api", return_value="Synthetic Python engineer") as call_ocr:
        kind, pages = server.extract_document("synthetic.png", payload)
        assert kind == "png" and pages == [(1, "Synthetic Python engineer")]
        call_ocr.assert_called_once()
        with Image.open(io.BytesIO(call_ocr.call_args.args[0])) as sent_image:
            assert sent_image.format == "PNG"
    with patch.object(server, "call_ocr_api", return_value="Synthetic JPEG engineer") as call_ocr:
        kind, pages = server.extract_document("synthetic.jpg", jpeg_bytes.getvalue())
        assert kind == "jpg" and pages == [(1, "Synthetic JPEG engineer")]
        call_ocr.assert_called_once()
    with patch.object(server, "call_ocr_api") as call_ocr:
        try:
            server.extract_document("mislabelled.jpg", payload)
        except ValueError as error:
            assert "tidak cocok dengan ekstensi" in str(error)
        else:
            raise AssertionError("image bytes that disagree with the filename extension should be rejected")
        call_ocr.assert_not_called()
    with patch.object(server, "MAX_IMAGE_PIXELS", 5):
        try:
            server.normalize_image(payload, expected_format="PNG")
        except ValueError as error:
            assert "Resolusi gambar" in str(error)
        else:
            raise AssertionError("image over the configured pixel limit should be rejected")


def make_synthetic_text_pdf(text: str, *, compressed: bool = False) -> bytes:
    import zlib

    escaped = text.encode("ascii").replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")
    stream = b"BT /F1 12 Tf 72 720 Td (" + escaped + b") Tj ET\n"
    stream_filter = b""
    if compressed:
        stream = zlib.compress(stream)
        stream_filter = b" /Filter /FlateDecode"
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        b"<< /Length " + str(len(stream)).encode("ascii") + stream_filter + b" >>\nstream\n" + stream + b"endstream",
    )
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(pdf)


def make_synthetic_image_pdf() -> bytes:
    import zlib

    image_data = zlib.compress(bytes((255, 255, 255, 40, 90, 140) * 3))
    content = b"q 3 0 0 2 0 0 cm /Im1 Do Q\n"
    image_object = (
        b"<< /Type /XObject /Subtype /Image /Width 3 /Height 2 /ColorSpace /DeviceRGB "
        b"/BitsPerComponent 8 /Filter /FlateDecode /Length "
        + str(len(image_data)).encode("ascii")
        + b" >>\nstream\n"
        + image_data
        + b"\nendstream"
    )
    objects = (
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        b"/Resources << /XObject << /Im1 4 0 R >> >> /Contents 5 0 R >>",
        image_object,
        b"<< /Length " + str(len(content)).encode("ascii") + b" >>\nstream\n" + content + b"endstream",
    )
    pdf = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, obj in enumerate(objects, start=1):
        offsets.append(len(pdf))
        pdf.extend(f"{number} 0 obj\n".encode("ascii") + obj + b"\nendobj\n")
    xref_offset = len(pdf)
    pdf.extend(f"xref\n0 {len(offsets)}\n".encode("ascii"))
    pdf.extend(b"0000000000 65535 f \n")
    for offset in offsets[1:]:
        pdf.extend(f"{offset:010d} 00000 n \n".encode("ascii"))
    pdf.extend(
        f"trailer\n<< /Size {len(offsets)} /Root 1 0 R >>\nstartxref\n{xref_offset}\n%%EOF\n".encode("ascii")
    )
    return bytes(pdf)


def check_text_document_extraction() -> None:
    kind, pages = server.extract_document("synthetic.txt", b"Python engineer.\nSQL.")
    assert kind == "txt" and pages == [(1, "Python engineer."), (1, "SQL.")]

    pdf = make_synthetic_text_pdf("Python engineer synthetic")
    kind, pages = server.extract_document("synthetic.pdf", pdf)
    assert kind == "pdf" and pages == [(1, "Python engineer synthetic")]

    from docx import Document

    document = Document()
    document.add_paragraph("Synthetic opening marker")
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Django"
    table.cell(0, 1).text = "PostgreSQL"
    document.add_paragraph("Python engineer synthetic")
    package = io.BytesIO()
    document.save(package)
    kind, pages = server.extract_document("synthetic.docx", package.getvalue())
    extracted = [text for _, text in pages]
    assert kind == "docx"
    assert extracted == [
        "Synthetic opening marker",
        "Django | PostgreSQL",
        "Python engineer synthetic",
    ], "DOCX paragraph and table order should be preserved"

    nested_document = Document()
    nested_cell = nested_document.add_table(rows=1, cols=1).cell(0, 0)
    nested_table = nested_cell.add_table(rows=1, cols=2)
    nested_table.cell(0, 0).text = "Python"
    nested_table.cell(0, 1).text = "FastAPI"
    nested_package = io.BytesIO()
    nested_document.save(nested_package)
    _, nested_pages = server.extract_document("nested-synthetic.docx", nested_package.getvalue())
    assert [text for _, text in nested_pages] == ["Python | FastAPI"]

    try:
        server.extract_document("broken.pdf", b"%PDF-1.4\n%%EOF\n")
    except ValueError as error:
        assert str(error) == "PDF rusak atau tidak dapat dibaca."
    else:
        raise AssertionError("malformed PDF should be reported as a safe document validation error")

    malformed_docx = io.BytesIO()
    with zipfile.ZipFile(malformed_docx, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<w:document")
    try:
        server.extract_document("broken.docx", malformed_docx.getvalue())
    except ValueError as error:
        assert str(error) == "DOCX rusak atau tidak dapat dibaca."
    else:
        raise AssertionError("malformed DOCX should be reported as a safe document validation error")


def check_document_resource_limits() -> None:
    package = io.BytesIO()
    with zipfile.ZipFile(package, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "x" * 128)
    docx_payload = package.getvalue()
    server.validate_docx_package(docx_payload)
    with patch.object(server, "MAX_DOCX_UNCOMPRESSED_BYTES", 32):
        try:
            server.validate_docx_package(docx_payload)
        except ValueError as error:
            assert "batas ukuran hasil ekstraksi" in str(error)
        else:
            raise AssertionError("DOCX package over the expanded-size limit should be rejected")
    with patch.object(server, "MAX_DOCX_ZIP_MEMBERS", 1):
        try:
            server.validate_docx_package(docx_payload)
        except ValueError as error:
            assert "komponen" in str(error)
        else:
            raise AssertionError("DOCX package over the component limit should be rejected")

    try:
        server.extract_document("empty.txt", b" \n\t")
    except ValueError as error:
        assert "tidak berisi teks" in str(error)
    else:
        raise AssertionError("empty text documents should be rejected")

    with patch.object(server, "MAX_EXTRACTED_TEXT_CHARACTERS", 4):
        try:
            server.extract_document("oversized-synthetic.txt", b"Python")
        except ValueError as error:
            assert "melampaui batas teks hasil ekstraksi" in str(error)
        else:
            raise AssertionError("documents over the extracted-character limit should be rejected")

        try:
            server.extract_document("oversized-synthetic.pdf", make_synthetic_text_pdf("Python"))
        except ValueError as error:
            assert "melampaui batas teks hasil ekstraksi" in str(error)
        else:
            raise AssertionError("PDF text over the extracted-character limit should be rejected")

    with patch.object(server, "MAX_PDF_DECODED_STREAM_BYTES", 64):
        try:
            server.extract_document(
                "oversized-decoded-stream.pdf",
                make_synthetic_text_pdf("A" * 256, compressed=True),
            )
        except ValueError as error:
            assert "PDF" in str(error) or "Halaman 1" in str(error)
        else:
            raise AssertionError("PDF decoded content over the worker stream limit should be rejected")

    with patch.object(server, "MAX_EXTRACTED_NONEMPTY_LINES", 2):
        _, accepted_pages = server.extract_document("line-boundary.txt", b"one\r\ntwo\n")
        assert accepted_pages == [(1, "one"), (1, "two")]
        try:
            server.extract_document("too-many-lines.txt", b"one\ntwo\nthree")
        except ValueError as error:
            assert "melampaui batas teks hasil ekstraksi" in str(error)
        else:
            raise AssertionError("documents over the nonempty-line limit should be rejected")

    with patch.object(server, "MAX_EXTRACTED_NONEMPTY_LINES", 2):
        pdf = make_synthetic_text_pdf("one\ntwo\nthree")
        try:
            server.extract_document("too-many-lines.pdf", pdf)
        except ValueError as error:
            assert "melampaui batas teks hasil ekstraksi" in str(error)
        else:
            raise AssertionError("PDF extraction should enforce the shared nonempty-line limit")

    from docx import Document

    document = Document()
    document.add_paragraph("one")
    document.add_paragraph("two")
    document.add_paragraph("three")
    package = io.BytesIO()
    document.save(package)
    with patch.object(server, "MAX_EXTRACTED_NONEMPTY_LINES", 2):
        try:
            server.extract_document("too-many-lines.docx", package.getvalue())
        except ValueError as error:
            assert "melampaui batas teks hasil ekstraksi" in str(error)
        else:
            raise AssertionError("DOCX extraction should enforce the shared nonempty-line limit")

    image_bytes = io.BytesIO()
    Image.new("RGB", (3, 2), color=(42, 91, 66)).save(image_bytes, format="PNG")
    with patch.object(server, "MAX_EXTRACTED_NONEMPTY_LINES", 2), patch.object(
        server, "call_ocr_api", return_value="one\ntwo\nthree"
    ):
        try:
            server.extract_document("too-many-lines.png", image_bytes.getvalue())
        except ValueError as error:
            assert "melampaui batas teks hasil ekstraksi" in str(error)
        else:
            raise AssertionError("OCR output should enforce the shared nonempty-line limit")


def check_pdf_worker_process_bounds() -> None:
    pdf = make_synthetic_text_pdf("worker timeout probe")
    spawned: list[subprocess.Popen[bytes]] = []
    worker_environments: list[dict[str, str] | None] = []
    real_popen = subprocess.Popen

    def record_process(*args, **kwargs):
        worker_environments.append(kwargs.get("env"))
        process = real_popen(*args, **kwargs)
        spawned.append(process)
        return process

    synthetic_secrets = {
        server.OCR_KEY_ENV: "synthetic-worker-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid",
        server.OCR_ALLOWED_ORIGINS_ENV: "https://ocr-stub.invalid",
        "HTTPS_PROXY": "https://proxy-stub.invalid",
    }
    with patch.dict(os.environ, synthetic_secrets), patch.object(
        server, "MAX_PDF_WORKER_SECONDS", -1
    ), patch.object(server.subprocess, "Popen", side_effect=record_process):
        try:
            server._run_pdf_parser_worker(pdf)
        except ValueError as error:
            assert "batas waktu" in str(error)
        else:
            raise AssertionError("PDF parser exceeding its time budget should be stopped")
    assert len(spawned) == 1 and spawned[0].poll() is not None, (
        "a timed-out PDF worker should be terminated and reaped"
    )
    assert len(worker_environments) == 1 and worker_environments[0] is not None
    assert not any(name in worker_environments[0] for name in synthetic_secrets), (
        "PDF workers should not inherit OCR credentials, endpoints, or proxy settings"
    )

    class OversizedWorkerResult:
        returncode = 0

        def communicate(self, _payload, timeout):
            return b"123456789", None

    with patch.object(server, "MAX_PDF_WORKER_RESULT_BYTES", 8), patch.object(
        server, "_assign_pdf_worker_job", return_value=None
    ), patch.object(server.subprocess, "Popen", return_value=OversizedWorkerResult()):
        try:
            server._run_pdf_parser_worker(pdf)
        except ValueError as error:
            assert "sumber daya parser" in str(error)
        else:
            raise AssertionError("oversized PDF worker output should be rejected")

    if os.name == "nt":
        process = subprocess.Popen(
            [sys.executable, "-c", "import time; time.sleep(60)"],
            stdin=subprocess.DEVNULL,
            stdout=subprocess.DEVNULL,
            stderr=subprocess.DEVNULL,
            creationflags=subprocess.CREATE_NO_WINDOW,
        )
        job_handle = None
        try:
            job_handle = server._assign_pdf_worker_job(process)
            assert job_handle is not None, "Windows PDF worker should receive a Job Object"
            server._close_pdf_worker_job(job_handle)
            job_handle = None
            assert process.wait(timeout=5) is not None, (
                "closing the PDF worker Job Object should terminate its process"
            )
        finally:
            if job_handle is not None:
                server._close_pdf_worker_job(job_handle)
            if process.poll() is None:
                process.kill()
                process.wait(timeout=5)
    print("PASS: PDF worker timeout, output-size, and platform containment checks")


def check_evidence_redaction() -> None:
    synthetic_line = (
        "Python engineer contact qa@example.com phone +1 555 010 2020 "
        "date of birth: 1994-03-12"
    )
    redacted = server.redact_for_evidence(synthetic_line)
    assert "qa@example.com" not in redacted
    assert "555 010 2020" not in redacted
    assert "1994-03-12" not in redacted
    assert "[email removed]" in redacted and "[phone removed]" in redacted

    indonesian_labeled_pii = "\n".join((
        "Nama lengkap: Kandidat Sintetis",
        "Email: kandidat.sintetis@example.test",
        "No. HP: +62 812-3456-7890",
        "Tempat, Tanggal Lahir: Jakarta, 12 Mei 1994",
        "Alamat domisili: Jalan Contoh 17, Jakarta",
        "Jenis Kelamin: Perempuan",
        "Agama: Islam",
        "Usia: 32",
        "NIK: 3174010212900001",
        "Paspor: X1234567",
        "Disabilitas: tidak ada",
        "Python services delivered for the synthetic product team.",
    ))
    indonesian_redacted = server.redact_for_evidence(indonesian_labeled_pii)
    for private_value in (
        "Kandidat Sintetis", "kandidat.sintetis@example.test", "812-3456-7890",
        "Jakarta", "12 Mei 1994", "Jalan Contoh", "Perempuan", "Islam", "Usia: 32",
        "3174010212900001", "X1234567", "Disabilitas",
    ):
        assert private_value.casefold() not in indonesian_redacted.casefold(), (
            f"Indonesian labeled personal data remains in evidence: {private_value}"
        )
    inline_age_text = server.redact_for_evidence(
        "Python services delivered; age: 32; berusia 31 tahun; 30 years old."
    )
    assert all(age not in inline_age_text for age in ("age: 32", "berusia 31 tahun", "30 years old"))
    assert inline_age_text.count("[age removed]") == 3
    experience_duration = server.redact_for_evidence(
        "Python engineer with 12 years of experience."
    )
    assert "12 years of experience" in experience_duration
    assert "Python services delivered" in indonesian_redacted

    manifest = json.loads((ROOT / "data" / "synthetic-cv-32" / "manifest.json").read_text(encoding="utf-8"))
    criteria = [
        {"id": f"smoke-{index}", "label": skill, "type": "required", "weight": 1.0}
        for index, skill in enumerate(server.SKILLS)
    ]
    leaked_samples = set()
    direct_pii_leaks = {key: set() for key in ("name", "address", "email", "phone", "dob")}
    evidence_count = 0
    for entry in manifest:
        sample_id = entry["sample_id"]
        sample_dir = ROOT / "data" / "synthetic-cv-32" / sample_id
        path = sample_dir / "resume_text.txt"
        expected = json.loads((sample_dir / "expected.json").read_text(encoding="utf-8"))
        names = [str(expected.get(key, "")).strip() for key in ("first_name", "last_name") if expected.get(key)]
        email = str(expected.get("email", "")).strip().casefold()
        phone_digits = re.sub(r"\D", "", str(expected.get("phone", "")))
        dob = str(expected.get("date_of_birth", "")).strip().casefold()
        address = str(expected.get("address", "")).strip().casefold()
        pages = [(1, line) for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
        _, rows = server.score_candidate(criteria, pages)
        for row in rows:
            if row["result"] == "unknown" or not row["snippet"]:
                continue
            evidence_count += 1
            snippet = row["snippet"]
            phone_leak = any(
                len(re.sub(r"\D", "", match.group(0))) >= 10
                for match in server.PHONE_RE.finditer(snippet)
            )
            if (server.EMAIL_RE.search(snippet) or phone_leak
                    or server.SENSITIVE_LINE_RE.match(snippet)
                    or re.search(r"\b(?:DOB|date of birth|birthday)\b", snippet, re.IGNORECASE)):
                leaked_samples.add(sample_id)
            folded = snippet.casefold()
            if len(names) == 2 and all(
                re.search(rf"(?<!\w){re.escape(name.casefold())}(?!\w)", folded)
                for name in names
            ):
                direct_pii_leaks["name"].add(sample_id)
            if email and email in folded:
                direct_pii_leaks["email"].add(sample_id)
            if phone_digits and phone_digits in re.sub(r"\D", "", snippet):
                direct_pii_leaks["phone"].add(sample_id)
            if dob and dob in folded:
                direct_pii_leaks["dob"].add(sample_id)
            if len(address) >= 5 and address in folded:
                direct_pii_leaks["address"].add(sample_id)
    assert evidence_count > 0, "synthetic evidence redaction check did not exercise any evidence rows"
    assert not leaked_samples, f"contact-like data remains in evidence for synthetic samples: {sorted(leaked_samples)}"
    assert not any(direct_pii_leaks.values()), (
        "direct PII from synthetic labels remains in evidence sample IDs: "
        f"{ {key: sorted(value) for key, value in direct_pii_leaks.items() if value} }"
    )


def check_negated_skill_evidence() -> None:
    criterion = {"id": "crit_python", "label": "Python", "type": "required", "weight": 2.0}
    negative_cases = (
        "No experience with Python.",
        "Not familiar with Python.",
        "Never used Python professionally.",
        "Without prior knowledge of Python.",
        "I don't know Python.",
        "Tidak berpengalaman menggunakan Python.",
        "Belum pernah bekerja dengan Python.",
        "Tanpa pengalaman Python.",
        "Kurang menguasai Python.",
        "Jarang menggunakan Python.",
    )
    for phrase in negative_cases:
        match = server.match_criterion(criterion, [(1, phrase)])
        assert match["result"] == "needs_verification", f"explicit negative should require verification: {phrase}"
        assert match["snippet"] and match["page_number"] == 1
        profile = server.profile_from_text(phrase)
        assert "python" not in profile["skills"], f"negative-only skill should not appear in profile: {phrase}"

    positive_phrase = "Built production APIs with Python."
    positive = server.match_criterion(criterion, [(2, positive_phrase)])
    assert positive["result"] == "matched" and positive["confidence"] == 1.0
    assert positive["snippet"] == positive_phrase and positive["page_number"] == 2
    positive_only = server.match_criterion(
        criterion, [(1, "Not only used Python but led its migration.")]
    )
    assert positive_only["result"] == "matched", "not-only construction should remain positive"
    assert server.match_criterion(criterion, [(1, "Bukan hanya menggunakan Python, tetapi juga membimbing tim.")])["result"] == "matched"
    assert server.match_criterion(criterion, [(1, "Kurang lebih 5 tahun menggunakan Python.")])["result"] == "matched"
    assert server.match_criterion(criterion, [(1, "Built production APIs with Python.")])["result"] == "matched"
    partial_phrase = "Key stakeholder led the department's management transition."
    partial_criterion = {
        "id": "crit_stakeholder", "label": "stakeholder management", "type": "required", "weight": 2.0,
    }
    partial = server.match_criterion(partial_criterion, [(3, partial_phrase)])
    assert partial["result"] == "partial" and partial["confidence"] == 0.5
    assert partial["snippet"] == partial_phrase and partial["page_number"] == 3
    partial_score, _ = server.score_candidate([partial_criterion], [(3, partial_phrase)])
    assert partial_score == 50.0
    unknown = server.match_criterion(criterion, [(4, "Experienced in Java and SQL.")])
    assert unknown["result"] == "unknown" and unknown["confidence"] == 0.0
    assert unknown["snippet"] == "" and unknown["page_number"] is None
    conflict = server.match_criterion(
        criterion,
        [(1, "No experience with Python."), (2, "Built production APIs with Python.")],
    )
    assert conflict["result"] == "needs_verification", "contradictory positive/negative evidence should be surfaced"
    assert conflict["snippet"] == "No experience with Python." and conflict["page_number"] == 1
    score, evidence = server.score_candidate([criterion], [(1, "No experience with Python.")])
    assert score == 0.0 and evidence[0]["result"] == "needs_verification"

    sensitive_only = server.profile_from_text("Name: Python\nAge: 30")
    assert "python" not in sensitive_only["skills"], "profile parser should ignore labeled sensitive-attribute lines"
    merged_section = server.profile_from_text("DOB: 20 Jan 1992 Skills Python")
    assert "python" in merged_section["skills"], "profile parser should retain a section merged after labeled PII"
    redacted_evidence = server.match_criterion(criterion, [(7, "Name: Synthetic Candidate; Python engineer")])
    assert redacted_evidence["result"] == "needs_verification" and redacted_evidence["snippet"] == ""
    redacted_score, _ = server.score_candidate([criterion], [(7, "Name: Synthetic Candidate; Python engineer")])
    assert redacted_score == 0.0, "evidence removed by redaction must not contribute a positive score"

    long_positive_line = "Early project details. " + ("routine work continued. " * 80) + "Built production APIs with Python."
    long_positive = server.match_criterion(criterion, [(4, long_positive_line)])
    assert long_positive["result"] == "matched"
    assert "Python" in long_positive["snippet"] and len(long_positive["snippet"]) <= 700

    long_negative_line = ("Earlier role details. " * 80) + "No experience with Python."
    long_negative = server.match_criterion(criterion, [(5, long_negative_line)])
    assert long_negative["result"] == "needs_verification"
    assert "No experience" in long_negative["snippet"] and "Python" in long_negative["snippet"]
    assert len(long_negative["snippet"]) <= 700

    long_partial_line = "stakeholder " + ("project delivery. " * 80) + "management"
    long_partial = server.match_criterion(partial_criterion, [(6, long_partial_line)])
    assert long_partial["result"] == "partial"
    assert all(term in long_partial["snippet"] for term in ("stakeholder", "management"))
    assert len(long_partial["snippet"]) <= 700


def check_candidate_score_ordering() -> None:
    result = synthetic_formats.check_candidate_score_ordering()
    assert result["candidate_count"] == 6
    assert result["score_assignments_match"] is True
    assert result["ordered_scores_nonincreasing"] is True
    assert result["timestamp_and_id_ties_deterministic"] is True
    assert result["all_checks_passed"] is True


def check_experience_duration_parsing() -> None:
    cases = (
        ("Age: 30 years old. Python engineer with five years of experience.", 5, "five years"),
        ("Candidate is 30 years of age; has 3 years of Python experience.", 3, "3 years"),
        ("Umur: 32 tahun. Pengalaman kerja selama lima tahun dengan Python.", 5, "selama lima tahun"),
        ("Berpengalaman selama satu tahun mengelola layanan Python.", 1, "selama satu tahun"),
        ("Python engineer with 3 thn experience.", 3, "3 thn"),
        ("Python engineer with 1.5 years of production experience.", 1.5, "1.5 years"),
        ("Pengalaman kerja dengan Python selama 1,5 tahun.", 1.5, "selama 1,5 tahun"),
        ("Memiliki pengalaman kerja selama 12 tahun di bidang teknologi.", 12, "selama 12 tahun"),
        ("Dua puluh tahun pengalaman mengelola sistem Python.", 20, "dua puluh tahun"),
        ("Usia 32 tahun.", None, None),
        ("Saya berumur 30 tahun.", None, None),
        ("Umur saya: 32 tahun.", None, None),
        ("Usia 30,5 tahun.", None, None),
    )
    for text, expected_years, expected_phrase in cases:
        profile = server.profile_from_text(text)
        assert profile["experience_years_mentioned"] == expected_years, (
            f"duration years mismatch for synthetic phrase: {text}"
        )
        assert profile["experience_duration_text"] == expected_phrase, (
            f"duration text mismatch for synthetic phrase: {text}"
        )


def check_education_mention_parsing() -> None:
    cases = (
        ("Bachelor of Science in Computer Science", ["Bachelor of Science"]),
        ("Master of Business Administration (MBA)", ["Master of Business Administration", "MBA"]),
        ("BSc, then Professional Diploma", ["BSc", "Professional Diploma"]),
        ("Associate of Arts and Diploma of Higher Education", ["Associate of Arts", "Diploma of Higher Education"]),
        ("S1 Informatika; Magister Manajemen; Doktor", ["S1", "Magister", "Doktor"]),
        ("D3 Keperawatan dan Sarjana Terapan", ["D3", "Sarjana Terapan"]),
    )
    for text, expected in cases:
        actual = server.profile_from_text(text)["education_levels_mentioned"]
        assert actual == expected, f"credential wording mismatch for synthetic phrase: {text}: {actual}"

    non_education_mentions = (
        "Acted as scrum master for a small product team.",
        "Customer Support Associate — Brightfen Service Desk Group",
        "IT Security Associate — Ashcombe Network Guard",
        "Certificates\nAdvanced Mobile Architecture Certificate",
    )
    for text in non_education_mentions:
        actual = server.profile_from_text(text)["education_levels_mentioned"]
        assert actual == [], f"non-degree wording was labeled as education: {text}: {actual}"

    sensitive_only = server.profile_from_text("Name: Bachelor of Science\nAge: Master of Arts")
    assert sensitive_only["education_levels_mentioned"] == [], (
        "education mentions from sensitive-only lines should be filtered"
    )
    merged_heading = server.profile_from_text("DOB: 20 Jan 1992 Education Bachelor of Science")
    assert merged_heading["education_levels_mentioned"] == ["Bachelor of Science"], (
        "education heading merged into a sensitive line should retain the following credential"
    )

    dataset = ROOT / "data" / "synthetic-cv-32"
    manifest = json.loads((dataset / "manifest.json").read_text(encoding="utf-8"))
    checked_labels = 0
    detected_mentions = 0
    for entry in manifest:
        sample_id = entry["sample_id"]
        text = (dataset / sample_id / "resume_text.txt").read_text(encoding="utf-8")
        expected = json.loads((dataset / sample_id / "expected.json").read_text(encoding="utf-8"))
        detected = [server.normalize(label) for label in server.profile_from_text(text)["education_levels_mentioned"]]
        detected_mentions += len(detected)
        for education in expected.get("educations", []):
            label = server.normalize(str(education.get("degree", "")))
            if label:
                assert label in server.normalize(text), f"synthetic source label missing from text: {sample_id} {label}"
                assert any(label in mention for mention in detected), (
                    f"education label not detected in synthetic source: {sample_id} {label}; parsed={detected}"
                )
                checked_labels += 1
    assert detected_mentions == checked_labels, (
        f"unexpected extra education mentions in synthetic corpus: {detected_mentions - checked_labels}"
    )
    print(f"PASS: literal credential mentions and {checked_labels} synthetic education labels are detected without level mapping")


def check_request_log_redaction() -> None:
    handler = object.__new__(server.Handler)
    handler.path = "/api/candidates/cand_synthetic_private_id/reviews?email=hidden@example.test"
    handler.command = "POST"
    handler.request_id = "0123456789abcdef0123456789abcdef"
    handler.response_bytes = 2
    output = io.StringIO()
    with redirect_stdout(output):
        handler.log_message(
            '"%s" %s %s',
            f"POST {handler.path} HTTP/1.1",
            "201",
            "-",
        )
    logged = output.getvalue()
    record = json.loads(logged)
    assert record["event"] == "http_request" and record["severity"] == "INFO"
    assert record["method"] == "POST" and record["route"] == "/api/candidates/:id/reviews"
    assert record["status"] == 201 and record["response_bytes"] == 2
    assert record["request_id"] == handler.request_id
    assert "cand_synthetic_private_id" not in logged
    assert "hidden@example.test" not in logged and "?email" not in logged
    assert "127.0.0.1" not in logged, "request logs should not retain client IP addresses"


def check_scanned_pdf_ocr_gate() -> None:
    dataset_dir = ROOT / "data" / "synthetic-cv-32"
    manifest = json.loads((dataset_dir / "manifest.json").read_text(encoding="utf-8"))
    sample_id = manifest[0]["sample_id"]
    pdf_path = dataset_dir / sample_id / f"{sample_id}.pdf"
    env = {
        server.OCR_KEY_ENV: "",
        "RECRUITMENT_COPILOT_OCR_API_URL": "",
    }
    with patch.dict(os.environ, env):
        try:
            server.extract_document(pdf_path.name, pdf_path.read_bytes())
        except ValueError as error:
            assert "perlu OCR" in str(error), f"unexpected scan handling message: {error}"
        else:
            raise AssertionError("image-only synthetic PDF should require OCR")
    with patch.object(server, "MAX_PDF_PAGES", 1):
        try:
            server.extract_document(pdf_path.name, pdf_path.read_bytes())
        except ValueError as error:
            assert "melebihi batas 1 halaman" in str(error)
        else:
            raise AssertionError("PDF over the configured page limit should be rejected")


def check_scanned_pdf_ocr_stub() -> None:
    env = {
        server.OCR_KEY_ENV: "synthetic-only-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid",
        server.OCR_ALLOWED_ORIGINS_ENV: "https://ocr-stub.invalid",
    }
    with patch.dict(os.environ, env), patch.object(
        server, "call_ocr_api", return_value="Synthetic scanned PDF engineer"
    ) as mocked_ocr, patch.object(server, "reserve_ocr_calls") as mocked_reserve:
        kind, pages = server.extract_document("synthetic-scan.pdf", make_synthetic_image_pdf())
    assert kind == "pdf" and pages == [(1, "Synthetic scanned PDF engineer")]
    mocked_ocr.assert_called_once()
    assert mocked_ocr.call_args.kwargs.get("reservation_made") is True
    mocked_reserve.assert_called_once_with(1)
    print("PASS: image-only synthetic PDF passes a normalized image to the OCR stub")


def check_pdf_ocr_image_budget() -> None:
    sample_id = "r00003"
    pdf_path = ROOT / "data" / "synthetic-cv-32" / sample_id / f"{sample_id}.pdf"
    env = {
        server.OCR_KEY_ENV: "test-only-stub-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid",
        server.OCR_ALLOWED_ORIGINS_ENV: "https://ocr-stub.invalid",
    }
    with patch.dict(os.environ, env), patch.object(server, "MAX_OCR_IMAGES_PER_PDF", 1):
        with patch.object(server.urllib.request, "build_opener") as mocked_build_opener:
            try:
                server.extract_document(pdf_path.name, pdf_path.read_bytes())
            except ValueError as error:
                assert "melebihi batas 1 gambar" in str(error)
            else:
                raise AssertionError("PDF over the OCR image budget should be rejected")
    mocked_build_opener.assert_not_called()


def check_ocr_response_limit() -> None:
    class OversizedResponse:
        def __init__(self):
            self.read_sizes = []

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self, size=-1):
            self.read_sizes.append(size)
            return b"x" * min(size, 16)

    class StubOpener:
        def __init__(self, response):
            self.response = response
            self.request = None
            self.timeout = None

        def open(self, request, timeout):
            self.request = request
            self.timeout = timeout
            return self.response

    response = OversizedResponse()
    opener = StubOpener(response)
    env = {
        server.OCR_KEY_ENV: "test-only-stub-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid",
        server.OCR_ALLOWED_ORIGINS_ENV: "https://ocr-stub.invalid",
    }
    with patch.dict(os.environ, env), patch.object(server, "MAX_OCR_RESPONSE_BYTES", 32):
        with patch.object(server.urllib.request, "build_opener", return_value=opener) as mocked_build_opener:
            try:
                server.call_ocr_api(b"synthetic-image", reservation_made=True)
            except ValueError as error:
                assert "melebihi batas 1 MiB" in str(error)
            else:
                raise AssertionError("oversized OCR response should be rejected")
    assert response.read_sizes == [33, 17, 1], "OCR response reader should stop after detecting the first byte over the cap"
    opener_handlers = mocked_build_opener.call_args.args
    assert isinstance(opener_handlers[0], urllib.request.ProxyHandler)
    assert opener_handlers[0].proxies == {}, "OCR must not inherit ambient proxy configuration"
    assert isinstance(opener_handlers[1], server._RejectOCRRedirects)
    assert opener.timeout == 45 and opener.request.full_url.endswith("/chat/completions")


def check_ocr_redirect_block() -> None:
    handler = server._RejectOCRRedirects()
    opener = urllib.request.build_opener(handler)
    redirect_handlers = [item for item in opener.handlers if isinstance(item, urllib.request.HTTPRedirectHandler)]
    assert len(redirect_handlers) == 1 and redirect_handlers[0] is handler
    request = urllib.request.Request("https://ocr-stub.invalid/chat/completions")
    try:
        handler.redirect_request(
            request, None, 302, "Found", {"Location": "https://unexpected.invalid/collect"},
            "https://unexpected.invalid/collect",
        )
    except urllib.error.HTTPError as error:
        assert error.code == 302
        assert error.url == request.full_url
    else:
        raise AssertionError("OCR redirect handler should reject redirection without opening its target")


def check_ocr_transport_security() -> None:
    env = {
        server.OCR_KEY_ENV: "test-only-stub-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid",
        server.OCR_ALLOWED_ORIGINS_ENV: "https://ocr-stub.invalid",
    }
    unsafe_urls = (
        "http://ocr-stub.invalid",
        "https://user:password@ocr-stub.invalid",
        "https://ocr-stub.invalid?token=secret",
        "https://ocr-stub.invalid#fragment",
    )
    with patch.dict(os.environ, env), patch.object(server.urllib.request, "build_opener") as mocked_build_opener:
        for unsafe_url in unsafe_urls:
            with patch.dict(os.environ, {"RECRUITMENT_COPILOT_OCR_API_URL": unsafe_url}):
                try:
                    server.call_ocr_api(b"synthetic-image", reservation_made=True)
                except ValueError as error:
                    assert "URL OCR" in str(error)
                else:
                    raise AssertionError(f"unsafe OCR URL should be rejected: {unsafe_url}")
    mocked_build_opener.assert_not_called()


def check_ocr_origin_allowlist() -> None:
    base_env = {
        server.OCR_KEY_ENV: "test-only-stub-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid/v1",
    }
    rejected_origins = (
        "",
        "https://unexpected.invalid",
        "https://ocr-stub.invalid:444",
        "https://ocr-stub.invalid/v1",
        "http://ocr-stub.invalid",
        "https://*.invalid",
    )
    with patch.object(server.urllib.request, "build_opener") as mocked_build_opener:
        for origin in rejected_origins:
            with patch.dict(os.environ, {**base_env, server.OCR_ALLOWED_ORIGINS_ENV: origin}):
                assert not server.ocr_is_configured(), f"unapproved OCR origin should not be configured: {origin}"
                try:
                    server.call_ocr_api(b"synthetic-image", reservation_made=True)
                except ValueError:
                    pass
                else:
                    raise AssertionError(f"unapproved OCR origin should be rejected: {origin}")
        mocked_build_opener.assert_not_called()

    response_body = json.dumps({"choices": [{"message": {"content": "synthetic OCR text"}}]}).encode("utf-8")

    class StubResponse:
        def __init__(self):
            self.sent = False

        def __enter__(self):
            return self

        def __exit__(self, *_args):
            return False

        def read(self, _size=-1):
            if self.sent:
                return b""
            self.sent = True
            return response_body

    class StubOpener:
        def __init__(self):
            self.request = None

        def open(self, request, timeout):
            self.request = request
            assert timeout == 45
            return StubResponse()

    opener = StubOpener()
    allowed_env = {
        **base_env,
        server.OCR_ALLOWED_ORIGINS_ENV: "https://OCR-STUB.INVALID.:443",
    }
    with patch.dict(os.environ, allowed_env):
        assert server.ocr_is_configured(), "host case, trailing dot, and explicit default port should normalize"
        with patch.object(server.urllib.request, "build_opener", return_value=opener):
            text = server.call_ocr_api(b"synthetic-image", reservation_made=True)
    assert text == "synthetic OCR text"
    assert opener.request.full_url == "https://ocr-stub.invalid/v1/chat/completions"


def check_ocr_configuration_health(base: str) -> None:
    missing_config = {
        server.OCR_KEY_ENV: "",
        server.OCR_URL_ENV: "",
        server.OCR_ALLOWED_ORIGINS_ENV: "",
    }
    with patch.dict(os.environ, missing_config):
        status, health = request_json(f"{base}/api/health")
        assert status == 200 and health["ocr_configured"] is False

    valid_config = {
        server.OCR_KEY_ENV: "test-only-stub-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid/v1",
        server.OCR_ALLOWED_ORIGINS_ENV: "https://ocr-stub.invalid",
    }
    with patch.dict(os.environ, valid_config):
        status, health = request_json(f"{base}/api/health")
        assert status == 200 and health["ocr_configured"] is True
        serialized = json.dumps(health)
        assert "test-only-stub-key" not in serialized
        assert "ocr-stub.invalid" not in serialized


def check_ocr_evaluator_no_network() -> None:
    sample_id = "r00003"
    reference_path = ROOT / "data" / "synthetic-cv-32" / sample_id / "resume_text.txt"
    reference_lines = reference_path.read_text(encoding="utf-8").splitlines()
    midpoint = len(reference_lines) // 2
    fake_pages = [(1, "\n".join(reference_lines[:midpoint])), (2, "\n".join(reference_lines[midpoint:]))]
    env = {
        server.OCR_KEY_ENV: "test-only-stub-key",
        server.OCR_URL_ENV: "https://ocr-stub.invalid",
        server.OCR_ALLOWED_ORIGINS_ENV: "https://ocr-stub.invalid",
    }
    with patch.dict(os.environ, env):
        with patch.object(synthetic_ocr.server, "extract_document", return_value=("pdf", fake_pages)) as mocked:
            result = synthetic_ocr.evaluate(sample_id, send_to_internal_ocr=True)
    mocked.assert_called_once()
    assert result["network_mode"] == "internal_ocr_opt_in"
    assert result["ocr_requests_sent"] == 2
    assert result["ocr_pages_returned"] == 2
    assert result["word_error_rate"] == 0.0


def check_ocr_preflight_no_network() -> None:
    with patch.object(synthetic_ocr.server, "extract_document") as mocked_extract:
        result = synthetic_ocr.evaluate("r00003", send_to_internal_ocr=False)
    mocked_extract.assert_not_called()
    assert result["status"] == "preflight_only"
    assert result["network_mode"] == "disabled"
    assert result["ocr_requests_sent"] == 0
    assert result["pipeline_version"] == server.pipeline_version()
    assert result["ocr_model"] == server.OCR_MODEL
    for fingerprint in (
        "dataset_manifest_sha256",
        "sample_pdf_sha256",
        "reference_text_sha256",
        "evaluator_sha256",
    ):
        assert re.fullmatch(r"[a-f0-9]{64}", result[fingerprint])


def check_windows_local_ocr_wrapper() -> None:
    runtime_environment = synthetic_ocr.local_ocr_environment()
    assert server.OCR_KEY_ENV not in runtime_environment
    assert server.OCR_URL_ENV not in runtime_environment
    assert "PATH" in runtime_environment
    if os.name != "nt":
        print("SKIP: Windows OCR process wrapper is Windows-only")
        return

    with tempfile.TemporaryDirectory(prefix="karsahire-windows-ocr-") as directory:
        root = Path(directory)
        executable = root / "powershell.exe"
        bridge = root / "bridge.ps1"
        executable.write_bytes(b"synthetic executable placeholder")
        bridge.write_text("# synthetic bridge fixture", encoding="utf-8")
        observed: dict = {}

        def stub_run(args, **kwargs):
            observed["args"] = args
            observed.update(kwargs)
            return synthetic_ocr.subprocess.CompletedProcess(
                args, 0, stdout=b"Synthetic OCR Test 12345\n"
            )

        secret_environment = {
            server.OCR_KEY_ENV: "test-only-local-secret",
            server.OCR_URL_ENV: "https://ocr-stub.invalid",
        }
        with patch.dict(os.environ, secret_environment):
            with patch.object(synthetic_ocr.subprocess, "run", side_effect=stub_run) as mocked_run:
                recognized = synthetic_ocr.local_windows_ocr_text(
                    b"synthetic PNG bytes", "en-US", str(executable), bridge
                )
        mocked_run.assert_called_once()
        assert recognized == "Synthetic OCR Test 12345\n"
        assert observed["input"] == b"synthetic PNG bytes"
        assert observed["env"]["KARSAHIRE_WINDOWS_OCR_LANGUAGE"] == "en-US"
        assert server.OCR_KEY_ENV not in observed["env"]
        assert server.OCR_URL_ENV not in observed["env"]
        assert observed["timeout"] == 30 and observed["check"] is False
        assert observed["args"][1:3] == ["-NoProfile", "-NonInteractive"]
        assert observed.get("shell") is not True

    print("PASS: local Windows OCR wrapper uses bounded stdin, rejects credential inheritance, and emits no network request")


def check_synthetic_ocr_corpus_aggregation() -> None:
    manifest_bytes = synthetic_ocr.resolve_dataset_file(
        synthetic_ocr.DATASET / "manifest.json"
    ).read_bytes()
    manifest_sha256 = synthetic_ocr_corpus.sha256_bytes(manifest_bytes)
    fake_sample = {
        "dataset_manifest_sha256": manifest_sha256,
        "sample_pdf_sha256": "a" * 64,
        "reference_text_sha256": "b" * 64,
        "pdf_pages": 2,
        "ocr_images_planned": 1,
        "ocr_images_processed": 1,
        "ocr_requests_sent": 0,
        "network_mode": "disabled_local_windows_ocr",
        "reference_word_count": 10,
        "ocr_word_count": 10,
        "word_edits": 1,
        "word_error_rate": 0.1,
        "pipeline_version": "local-synthetic-test",
        "evaluator_sha256": "c" * 64,
        "ocr_bridge_sha256": "d" * 64,
        "ocr_engine_version": "Windows OCR synthetic stub",
    }
    with patch.object(synthetic_ocr_corpus.os, "name", "nt"):
        with patch.object(synthetic_ocr_corpus.evaluator, "evaluate", return_value=fake_sample) as mocked:
            with patch.object(synthetic_ocr_corpus.sys, "stderr", io.StringIO()):
                result = synthetic_ocr_corpus.evaluate_corpus("en-US")
    assert result["status"] == "complete"
    assert result["samples_attempted"] == result["samples_processed"] == 32
    assert result["samples_failed"] == 0 and result["network_mode"] == "disabled_local_windows_ocr"
    assert result["dataset_manifest_sha256"] == manifest_sha256
    assert result["ocr_images_processed"] == 32 and result["ocr_requests_sent"] == 0
    assert result["reference_word_count"] == 320 and result["word_edits"] == 32
    assert result["corpus_micro_word_error_rate"] == 0.1
    assert result["dataset_snapshot_sha256"] == synthetic_ocr_corpus.pinned_dataset_snapshot_sha256(
        synthetic_ocr.DATASET
    )
    assert re.fullmatch(r"[a-f0-9]{64}", result["runner_sha256"])
    assert re.fullmatch(r"[a-f0-9]{64}", result["dataset_fingerprint_helper_sha256"])
    assert mocked.call_count == 32
    assert all(call.kwargs.get("use_local_windows_ocr") is True for call in mocked.call_args_list)

    partial_sample = {**fake_sample, "ocr_images_processed": 0}
    with patch.object(synthetic_ocr_corpus.os, "name", "nt"):
        with patch.object(synthetic_ocr_corpus.evaluator, "evaluate", return_value=partial_sample):
            with patch.object(synthetic_ocr_corpus.sys, "stderr", io.StringIO()):
                partial_result = synthetic_ocr_corpus.evaluate_corpus("en-US")
    assert partial_result["status"] == "partial" and partial_result["samples_failed"] == 32
    assert partial_result["corpus_micro_word_error_rate"] is None
    assert partial_result["reference_word_count"] is None and partial_result["word_edits"] is None
    print("PASS: synthetic OCR corpus runner verifies pinned inputs, complete-image coverage, and zero network calls")


def check_synthetic_format_roundtrip() -> None:
    result = synthetic_formats.evaluate()
    assert result["samples_processed"] == 32
    assert result["pipeline_version"] == server.pipeline_version()
    assert result["formats"]["txt"]["text_match_count"] == 32
    assert result["formats"]["docx"]["text_match_count"] == 32
    assert result["formats"]["text_pdf"]["text_match_count"] == 32
    assert result["formats"]["text_pdf"]["profile_match_count"] == 32
    assert result["pdf_ascii_punctuation_probe"]["matched"] is True
    evidence_probe = result["evidence_excerpt_focus_probe"]
    assert evidence_probe["cases_checked"] == 7
    assert evidence_probe["positive_match_visible"] is True
    assert evidence_probe["negative_cue_visible"] is True
    assert evidence_probe["partial_terms_visible"] is True
    assert evidence_probe["page_attribution_preserved"] is True
    assert evidence_probe["all_snippets_within_limit"] is True
    assert evidence_probe["redacted_only_evidence_unscored"] is True
    assert evidence_probe["merged_sensitive_header_section_retained"] is True
    assert evidence_probe["sensitive_education_lines_filtered"] is True
    assert evidence_probe["merged_sensitive_education_heading_retained"] is True
    score_order_probe = result["candidate_score_order_probe"]
    assert score_order_probe["candidate_count"] == 6
    assert score_order_probe["score_assignments_match"] is True
    assert score_order_probe["ordered_scores_nonincreasing"] is True
    assert score_order_probe["timestamp_and_id_ties_deterministic"] is True
    assert result["pdf_page_attribution_probe"]["matched"] is True
    print("PASS: pinned local TXT/DOCX/text-PDF round-trip, ASCII apostrophe, and PDF page attribution (32 synthetic samples)")


def check_release_preflight_environment_filter() -> None:
    injected = {
        server.OCR_KEY_ENV: "synthetic-secret-that-must-not-propagate",
        server.OCR_URL_ENV: "https://ocr-stub.invalid",
        "KARSAHIRE_PREFLIGHT_TEST_SECRET": "never-forward-this",
    }
    with patch.dict(os.environ, injected):
        child_env = local_release.child_environment()
    assert set(child_env).issubset(set(local_release.SAFE_CHILD_ENV))
    assert server.OCR_KEY_ENV not in child_env and server.OCR_URL_ENV not in child_env
    assert "KARSAHIRE_PREFLIGHT_TEST_SECRET" not in child_env
    with tempfile.TemporaryDirectory(prefix="karsahire-preflight-report-") as directory:
        report_path = Path(directory) / "preflight.json"
        report_bytes = b'{"status":"local_preflight_passed"}\n'
        published = local_release.write_report_exclusive(report_path, report_bytes)
        assert published == report_path.resolve()
        assert report_path.read_bytes() == report_bytes
        assert sorted(path.name for path in report_path.parent.iterdir()) == ["preflight.json"]
        try:
            local_release.write_report_exclusive(report_path, b'{"status":"replacement"}\n')
        except FileExistsError:
            pass
        else:
            raise AssertionError("preflight report publication must refuse replacement")
        assert report_path.read_bytes() == report_bytes
        assert sorted(path.name for path in report_path.parent.iterdir()) == ["preflight.json"]
    print("PASS: local preflight filters secrets and atomically saves a report without overwriting")


def check_tessdata_fingerprinting() -> None:
    with tempfile.TemporaryDirectory(prefix="karsahire-tessdata-") as directory:
        prefix = Path(directory)
        data_dir = prefix / "tessdata"
        data_dir.mkdir()
        model = data_dir / "eng.traineddata"
        model.write_bytes(b"synthetic-traineddata-fixture")

        resolved_prefix, fingerprint = synthetic_ocr.resolve_tessdata_prefix(
            str(prefix / "bin" / "tesseract.exe"), "eng", str(prefix)
        )
        model_hash = synthetic_ocr.file_sha256(model).encode("ascii")
        expected = synthetic_ocr.hashlib.sha256(b"eng\0" + model_hash + b"\0").hexdigest()
        assert resolved_prefix == prefix.resolve()
        assert fingerprint == expected

        direct_prefix, direct_fingerprint = synthetic_ocr.resolve_tessdata_prefix(
            str(prefix / "bin" / "tesseract.exe"), "eng", str(data_dir)
        )
        assert direct_prefix == prefix.resolve() and direct_fingerprint == fingerprint
        try:
            synthetic_ocr.resolve_tessdata_prefix(
                str(prefix / "bin" / "tesseract.exe"), "ind", str(prefix)
            )
        except ValueError as error:
            assert "Data bahasa Tesseract 'ind' tidak ditemukan" in str(error)
        else:
            raise AssertionError("missing language model must not produce a fingerprint")
    print("PASS: synthetic Tesseract model fingerprints are stable and missing languages fail closed")


def check_synthetic_baseline_breakdown() -> None:
    result = evaluate_synthetic.evaluate()
    overlap = result["strict_reference_overlap"]
    breakdown = overlap["predicted_skills_without_matching_reference_label"]
    assert result["samples_processed"] == 32
    assert result["pipeline_version"].startswith("local-")
    assert result["network_mode"] == "disabled"
    for fingerprint in (
        "dataset_manifest_sha256",
        "dataset_content_sha256",
        "evaluator_sha256",
    ):
        assert re.fullmatch(r"[a-f0-9]{64}", result[fingerprint])
    assert sum(breakdown.values()) == overlap["prediction_without_reference_label_count"]
    assert overlap["prediction_without_reference_label_count"] == overlap["predictions_without_matching_reference_label"]
    assert "false_positive" not in overlap, "unmatched reference labels should not be presented as verified errors"
    assert "not verified errors" in overlap["interpretation"]
    assert all(isinstance(skill, str) and count > 0 for skill, count in breakdown.items())
    assert "r000" not in json.dumps(result), "baseline output should remain aggregate-only"
    with tempfile.TemporaryDirectory(prefix="karsahire-manifest-path-") as directory:
        dataset_root = Path(directory)
        sample_dir = dataset_root / "r00001"
        sample_dir.mkdir()
        (dataset_root / "manifest.json").write_text(
            json.dumps([{"sample_id": "r00001"}]), encoding="utf-8"
        )
        (sample_dir / "resume_text.txt").write_text("Worked with Python.", encoding="utf-8")
        expected_path = sample_dir / "expected.json"
        expected_path.write_text(json.dumps({"skills": []}), encoding="utf-8")
        (sample_dir / "metadata.json").write_text(json.dumps({"years_experience": None}), encoding="utf-8")
        first_fixture_result = evaluate_synthetic.evaluate(dataset_root)
        expected_path.write_text(
            json.dumps({"skills": [{"skill_name": "Python"}]}), encoding="utf-8"
        )
        changed_fixture_result = evaluate_synthetic.evaluate(dataset_root)
        assert first_fixture_result["dataset_manifest_sha256"] == changed_fixture_result["dataset_manifest_sha256"]
        assert first_fixture_result["dataset_content_sha256"] != changed_fixture_result["dataset_content_sha256"]
        (dataset_root / "manifest.json").write_text(
            json.dumps([{"sample_id": "../outside"}]), encoding="utf-8"
        )
        try:
            evaluate_synthetic.evaluate(dataset_root)
        except ValueError as error:
            assert "ID sampel" in str(error)
        else:
            raise AssertionError("synthetic evaluator should reject path traversal in manifest sample IDs")


def _relative_luminance(color: str) -> float:
    channels = [int(color[index:index + 2], 16) / 255 for index in (1, 3, 5)]
    linear = [value / 12.92 if value <= 0.04045 else ((value + 0.055) / 1.055) ** 2.4 for value in channels]
    return 0.2126 * linear[0] + 0.7152 * linear[1] + 0.0722 * linear[2]


def _contrast_ratio(left: str, right: str) -> float:
    a, b = _relative_luminance(left), _relative_luminance(right)
    return (max(a, b) + 0.05) / (min(a, b) + 0.05)


def check_chart_color_contrast() -> None:
    css = (ROOT / "web" / "app.css").read_text(encoding="utf-8")
    tokens = {
        name: value.lower()
        for name, value in re.findall(r"--([a-z0-9-]+)\s*:\s*(#[0-9a-fA-F]{6})", css)
    }

    track_match = re.search(r"\.evidence-bar\s*\{[^}]*?background:\s*(#[0-9a-fA-F]{6})", css, re.DOTALL)
    assert track_match, "evidence chart track color is missing"
    page_bg, track = tokens.get("bg"), track_match.group(1).lower()
    assert page_bg, "page background color token is missing"
    fills = {name: tokens.get(f"evidence-{name}") for name in ("found", "partial", "verify", "unknown")}
    for name, color in fills.items():
        assert color, f"chart color token evidence-{name} is missing"
        assert _contrast_ratio(color, page_bg) >= 4.5, (
            f"{name} status text contrast against the page background is below 4.5:1"
        )
        assert _contrast_ratio(color, track) >= 3.0, f"{name} chart fill contrast against its track is below 3:1"
        assert _contrast_ratio(color, "#ffffff") >= 3.0, f"{name} chart fill contrast against its divider is below 3:1"

    app_js = (ROOT / "web" / "app.js").read_text(encoding="utf-8")
    for label in ("Ada bukti tekstual", "Bukti parsial", "Perlu verifikasi", "Belum ditemukan"):
        assert label in app_js, f"chart status label is missing: {label}"
    assert "calculatePercentageShares(counts, total)" in app_js, "chart shares should use all evidence rows as the denominator"
    assert "percentageFormatter.format(shares[key])" in app_js, "chart status shares should be rounded consistently"
    assert "<strong>${counts[key]} <span class=\"legend-share\">${share}</span></strong>" in app_js, (
        "chart status counts and percentage shares should remain available as text"
    )


def check_reduced_motion_contract() -> None:
    css = (ROOT / "web" / "app.css").read_text(encoding="utf-8")
    match = re.search(
        r"@media\s*\(\s*prefers-reduced-motion\s*:\s*reduce\s*\)\s*\{(?P<body>[\s\S]*?)^\}",
        css,
        re.MULTILINE,
    )
    assert match, "reduced-motion media query is missing"
    body = match.group("body")
    required_rules = (
        (r"html:focus-within\s*\{[^}]*scroll-behavior\s*:\s*auto\s*!important", "smooth scrolling should be disabled"),
        (r"animation-duration\s*:\s*0\.01ms\s*!important", "animations should be minimized"),
        (r"animation-iteration-count\s*:\s*1\s*!important", "animations should not repeat"),
        (r"transition-duration\s*:\s*0\.01ms\s*!important", "transitions should be minimized"),
    )
    for pattern, message in required_rules:
        assert re.search(pattern, body, re.DOTALL), message


def check_reviewer_form_accessibility() -> None:
    app_js = (ROOT / "web" / "app.js").read_text(encoding="utf-8")
    css = (ROOT / "web" / "app.css").read_text(encoding="utf-8")

    note_markup = re.search(
        r'<label(?P<label_attrs>[^>]*)>Catatan berdasarkan kriteria dan bukti \(opsional\)</label>'
        r'\s*<textarea(?P<textarea_attrs>[^>]*)name="note"[^>]*>',
        app_js,
    )
    assert note_markup, "review note should have a programmatic label"
    label_for = re.search(r'\bfor="([^"]+)"', note_markup.group("label_attrs"))
    textarea_id = re.search(r'\bid="([^"]+)"', note_markup.group("textarea_attrs"))
    assert label_for and textarea_id and label_for.group(1) == textarea_id.group(1), (
        "review note label target should match the textarea ID"
    )
    assert '${escapeHtml(candidate.id)}' in textarea_id.group(1), (
        "review note IDs should stay unique to their candidate"
    )

    for selector in (
        ".review-form input:focus-visible",
        ".review-form select:focus-visible",
        ".review-form textarea:focus-visible",
    ):
        assert selector in css, f"review control needs visible keyboard focus: {selector}"
    focus_rule = re.search(r"\.review-form select:focus-visible[^{}]*\{(?P<body>[^{}]*)\}", css)
    assert focus_rule and re.search(r"outline\s*:\s*3px solid #335a48", focus_rule.group("body")), (
        "review controls should keep the strong visible focus outline"
    )


def check_annotation_evaluator() -> None:
    rows = [
        ("matched", "supported", "supported", "supported", "software_engineering", "en", "pdf_text"),
        ("partial", "partial", "no_evidence", "partial", "software_engineering", "en", "pdf_text"),
        ("unknown", "no_evidence", "no_evidence", "no_evidence", "software_engineering", "id", "pdf_scan"),
        ("matched", "no_evidence", "no_evidence", "no_evidence", "product_design", "id", "docx"),
        ("needs_verification", "needs_verification", "no_evidence", "needs_verification", "product_design", "en", "pdf_scan"),
        ("parse_error", "needs_verification", "no_evidence", "needs_verification", "software_engineering", "id", "pdf_scan"),
    ]
    reference_reviews = [
        ("valid", "valid", "valid"),
        ("missing", "missing", "missing"),
        ("not_applicable", "not_applicable", "not_applicable"),
        ("valid", "incorrect", "valid"),
        ("valid", "valid", "valid"),
        ("not_applicable", "not_applicable", "not_applicable"),
    ]
    with tempfile.TemporaryDirectory(prefix="karsahire-annotation-") as directory:
        csv_path = Path(directory) / "annotations.csv"
        with csv_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            for index, (model, recruiter, manager, adjudicated, job_family, cv_language, source_format) in enumerate(rows, start=1):
                recruiter_reference, manager_reference, adjudicated_reference = reference_reviews[index - 1]
                writer.writerow({
                    "requisition_id": "job_0123456789ab",
                    "candidate_id": f"cand_{index:012x}",
                    "criteria_version": "v1",
                    "criterion_id": "crit_0123456789ab",
                    "pipeline_version": "local-development",
                    "job_family": job_family,
                    "cv_language": cv_language,
                    "source_format": source_format,
                    "model_result": model,
                    "model_evidence_ref": "" if model in {"partial", "unknown", "parse_error"} else f"ev_{index:012x}",
                    "model_evidence_page": (
                        "1" if model == "matched" and source_format == "pdf_text"
                        else "2" if model == "needs_verification" and source_format == "pdf_scan"
                        else ""
                    ),
                    "recruiter_reference_label": recruiter_reference,
                    "hiring_manager_reference_label": manager_reference,
                    "adjudicated_reference_label": adjudicated_reference,
                    "recruiter_label": recruiter,
                    "hiring_manager_label": manager,
                    "adjudicated_label": adjudicated,
                })
        result = evaluate_adjudications.evaluate(csv_path)
        incomplete_reference_adjudication_rows = []
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            incomplete_reference_adjudication_rows = list(reader)
            fieldnames = reader.fieldnames
        incomplete_reference_adjudication_rows[0]["adjudicated_reference_label"] = ""
        incomplete_reference_adjudication_path = Path(directory) / "partial-reference-adjudication.csv"
        with incomplete_reference_adjudication_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(incomplete_reference_adjudication_rows)
        partial_reference_result = evaluate_adjudications.evaluate(incomplete_reference_adjudication_path)
        assert partial_reference_result["evidence_reference_review"]["review_coverage"]["adjudication_unreviewed_rows"] is None
        assert partial_reference_result["evidence_reference_review"]["review_coverage"]["adjudication_rate"] is None
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            unexpected_locator_rows = list(reader)
            fieldnames = reader.fieldnames
        unexpected_locator_rows[2]["model_evidence_ref"] = "ev_000000000003"
        unexpected_locator_rows[2]["model_evidence_page"] = "1"
        unexpected_locator_rows[2]["recruiter_reference_label"] = "incorrect"
        unexpected_locator_rows[2]["hiring_manager_reference_label"] = "incorrect"
        unexpected_locator_rows[2]["adjudicated_reference_label"] = "incorrect"
        unexpected_locator_path = Path(directory) / "unexpected-locator.csv"
        with unexpected_locator_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(unexpected_locator_rows)
        unexpected_locator_result = evaluate_adjudications.evaluate(unexpected_locator_path)
        unexpected_reference_metrics = unexpected_locator_result["evidence_reference_review"]["adjudicated"]
        assert unexpected_reference_metrics["status_counts"]["incorrect"] is None
        assert unexpected_reference_metrics["status_counts_suppressed"] is True
        assert unexpected_reference_metrics["citation_validity_rate"] is None
        assert unexpected_reference_metrics["expected_locator_rows"] is None
        unsafe_path = Path(directory) / "unsafe-ids.csv"
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            unsafe_rows = list(reader)
            fieldnames = reader.fieldnames
        unsafe_rows[0]["candidate_id"] = "person@example.test"
        with unsafe_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(unsafe_rows)
        try:
            evaluate_adjudications.evaluate(unsafe_path)
        except ValueError as error:
            assert "pseudonim berformat aman" in str(error)
            assert "person@example.test" not in str(error)
        else:
            raise AssertionError("annotation evaluator should reject identifiers that may contain PII")
        unsafe_stratum_path = Path(directory) / "unsafe-stratum.csv"
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            unsafe_stratum_rows = list(reader)
            fieldnames = reader.fieldnames
        unsafe_stratum_rows[0]["job_family"] = "person@example.test"
        with unsafe_stratum_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(unsafe_stratum_rows)
        try:
            evaluate_adjudications.evaluate(unsafe_stratum_path)
        except ValueError as error:
            assert "job_family" in str(error)
            assert "person@example.test" not in str(error)
        else:
            raise AssertionError("annotation evaluator should reject free-text stratum values")
        mixed_version_path = Path(directory) / "mixed-pipeline-versions.csv"
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            mixed_version_rows = list(reader)
            fieldnames = reader.fieldnames
        mixed_version_rows[0]["pipeline_version"] = "local-development-v2"
        with mixed_version_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(mixed_version_rows)
        try:
            evaluate_adjudications.evaluate(mixed_version_path)
        except ValueError as error:
            assert "satu pipeline_version" in str(error)
        else:
            raise AssertionError("annotation evaluator should reject mixed pipeline versions")
        mixed_criteria_path = Path(directory) / "mixed-criteria-versions.csv"
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            mixed_criteria_rows = list(reader)
            fieldnames = reader.fieldnames
        mixed_criteria_rows[0]["criteria_version"] = "v2"
        with mixed_criteria_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(mixed_criteria_rows)
        try:
            evaluate_adjudications.evaluate(mixed_criteria_path)
        except ValueError as error:
            assert "satu criteria_version" in str(error)
        else:
            raise AssertionError("annotation evaluator should reject mixed criteria versions")
        incomplete_review_path = Path(directory) / "incomplete-adjudication.csv"
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            incomplete_review_rows = list(reader)
            fieldnames = reader.fieldnames
        incomplete_review_rows[0]["hiring_manager_label"] = ""
        with incomplete_review_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(incomplete_review_rows)
        try:
            evaluate_adjudications.evaluate(incomplete_review_path)
        except ValueError as error:
            assert "label kedua reviewer lengkap" in str(error)
        else:
            raise AssertionError("annotation evaluator should reject adjudication before independent review is complete")
        incomplete_reference_review_path = Path(directory) / "incomplete-reference-adjudication.csv"
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            incomplete_reference_rows = list(reader)
            fieldnames = reader.fieldnames
        incomplete_reference_rows[0]["hiring_manager_reference_label"] = ""
        with incomplete_reference_review_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(incomplete_reference_rows)
        try:
            evaluate_adjudications.evaluate(incomplete_reference_review_path)
        except ValueError as error:
            assert "Referensi pada baris" in str(error)
        else:
            raise AssertionError("annotation evaluator should reject citation adjudication before independent reference review is complete")
        missing_result_path = Path(directory) / "missing-model-result.csv"
        with csv_path.open("r", encoding="utf-8", newline="") as source:
            reader = csv.DictReader(source)
            missing_result_rows = list(reader)
            fieldnames = reader.fieldnames
        missing_result_rows[0]["model_result"] = ""
        with missing_result_path.open("w", encoding="utf-8", newline="") as target:
            writer = csv.DictWriter(target, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(missing_result_rows)
        try:
            evaluate_adjudications.evaluate(missing_result_path)
        except ValueError as error:
            assert "Label model tidak valid" in str(error)
        else:
            raise AssertionError("annotation evaluator should reject missing model outcomes")
        eligible_path = Path(directory) / "eligible-stratum.csv"
        eligible_rows = []
        with eligible_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            for index in range(100, 105):
                row = {
                    "requisition_id": "job_0123456789ab",
                    "candidate_id": f"cand_{index:012x}",
                    "criteria_version": "v1",
                    "criterion_id": "crit_0123456789ab",
                    "pipeline_version": "local-development",
                    "job_family": "software_engineering",
                    "cv_language": "en",
                    "source_format": "pdf_text",
                    "model_result": "matched",
                    "model_evidence_ref": f"ev_{index:012x}",
                    "model_evidence_page": "1",
                    "recruiter_reference_label": "valid",
                    "hiring_manager_reference_label": "valid",
                    "adjudicated_reference_label": "valid",
                    "recruiter_label": "supported",
                    "hiring_manager_label": "supported",
                    "adjudicated_label": "supported",
                }
                eligible_rows.append(row)
                writer.writerow(row)
        eligible_result = evaluate_adjudications.evaluate(eligible_path)
        incorrect_reference_path = Path(directory) / "incorrect-reference-eligible.csv"
        incorrect_reference_rows = [dict(row) for row in eligible_rows]
        for row in incorrect_reference_rows:
            row["recruiter_reference_label"] = "incorrect"
            row["hiring_manager_reference_label"] = "incorrect"
            row["adjudicated_reference_label"] = "incorrect"
        with incorrect_reference_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            writer.writerows(incorrect_reference_rows)
        incorrect_reference_result = evaluate_adjudications.evaluate(incorrect_reference_path)
        single_sparse_outcome_path = Path(directory) / "single-sparse-outcome.csv"
        single_sparse_outcome_rows = [dict(row) for row in eligible_rows]
        additional_row = dict(eligible_rows[-1])
        additional_row["candidate_id"] = "cand_0000000000aa"
        additional_row["model_evidence_ref"] = "ev_0000000000aa"
        single_sparse_outcome_rows.append(additional_row)
        for field in (
            "model_evidence_ref", "model_evidence_page",
            "recruiter_reference_label", "hiring_manager_reference_label",
            "adjudicated_reference_label",
        ):
            single_sparse_outcome_rows[0][field] = ""
        single_sparse_outcome_rows[0]["recruiter_reference_label"] = "missing"
        single_sparse_outcome_rows[0]["hiring_manager_reference_label"] = "missing"
        single_sparse_outcome_rows[0]["adjudicated_reference_label"] = "missing"
        with single_sparse_outcome_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            writer.writerows(single_sparse_outcome_rows)
        single_sparse_outcome_result = evaluate_adjudications.evaluate(single_sparse_outcome_path)
        partial_metrics_path = Path(directory) / "partial-metric-participation.csv"
        partial_metrics_rows = [dict(row) for row in eligible_rows]
        for row in partial_metrics_rows[1:]:
            for field in (
                "recruiter_label", "hiring_manager_label", "adjudicated_label",
                "recruiter_reference_label", "hiring_manager_reference_label",
                "adjudicated_reference_label",
            ):
                row[field] = ""
        with partial_metrics_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            writer.writerows(partial_metrics_rows)
        partial_metrics_result = evaluate_adjudications.evaluate(partial_metrics_path)
        sparse_path = Path(directory) / "complementary-suppression.csv"
        sparse_rows = [dict(row) for row in eligible_rows]
        sparse_row = dict(eligible_rows[-1])
        sparse_row["candidate_id"] = "cand_000000000069"
        sparse_row["model_evidence_ref"] = "ev_000000000069"
        sparse_row["job_family"] = "product_design"
        sparse_rows.append(sparse_row)
        with sparse_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            writer.writerows(sparse_rows)
        sparse_result = evaluate_adjudications.evaluate(sparse_path)
        repeated_candidate_path = Path(directory) / "same-candidate-across-requisitions.csv"
        with repeated_candidate_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            for index in range(5):
                row = dict(eligible_rows[0])
                row["requisition_id"] = f"job_{index + 300:012x}"
                row["candidate_id"] = "cand_000000000099"
                row["model_evidence_ref"] = f"ev_{index + 300:012x}"
                writer.writerow(row)
        repeated_candidate_result = evaluate_adjudications.evaluate(repeated_candidate_path)
    assert result["reviewer_agreement"]["double_reviewed_rows"] == 6
    assert result["criteria_version"] == "v1"
    assert result["pipeline_version"] == "local-development"
    assert re.fullmatch(r"[a-f0-9]{64}", result["evaluator_sha256"])
    assert result["overall_privacy"]["suppressed"] is True
    assert result["overall_privacy"]["sample_suppressed"] is False
    assert result["overall_privacy"]["distinct_candidate_count"] == 6
    assert "reviewer_agreement.outcome" in result["overall_privacy"]["suppressed_measures"]
    assert "model_vs_adjudication.accuracy_outcome" in result["overall_privacy"]["suppressed_measures"]
    assert "evidence_reference_review.adjudicated.status_counts" in result["overall_privacy"]["suppressed_measures"]
    assert result["reviewer_agreement"]["exact_agreement_count"] is None
    assert result["reviewer_agreement"]["exact_agreement_rate"] is None
    assert result["reviewer_agreement"]["agreement_suppressed"] is True
    assert result["model_vs_adjudication"]["compared_rows"] == 6
    assert result["model_vs_adjudication"]["accuracy"] is None
    assert result["model_vs_adjudication"]["accuracy_suppressed"] is True
    assert all(value is None for value in result["model_vs_adjudication"]["by_evidence_label"].values())
    reference_presence = result["evidence_locator_presence"]["reference"]
    assert reference_presence["suppressed"] is True
    assert reference_presence["expected_rows"] is None
    page_presence = result["evidence_locator_presence"]["page_number"]
    assert page_presence["suppressed"] is True
    assert page_presence["expected_rows"] is None
    reference_review = result["evidence_reference_review"]
    assert reference_review["review_coverage"]["input_rows"] == 6
    assert reference_review["review_coverage"]["reviewer_review_rate"] == 1.0
    assert reference_review["review_coverage"]["adjudication_unreviewed_rows"] == 0
    assert reference_review["review_coverage"]["adjudication_rate"] == 1.0
    assert reference_review["reviewer_double_reviewed_rows"] == 6
    assert reference_review["reviewer_agreement_rows"] is None
    assert reference_review["reviewer_agreement"] is None
    assert reference_review["adjudicated"]["status_counts"] == {
        "valid": None, "incorrect": 0, "missing": None, "not_applicable": None,
    }
    assert reference_review["adjudicated"]["status_counts_suppressed"] is True
    assert reference_review["adjudicated"]["citation_validity_rate"] is None
    assert reference_review["adjudicated"]["citation_coverage_rate"] is None
    assert reference_review["reviewer_agreement_by_stratum"]["source_format"] == {"reported": {}, "suppressed": True}
    assert reference_review["adjudicated_by_stratum"]["job_family"] == {"reported": {}, "suppressed": True}
    assert reference_presence["by_stratum"]["source_format"] == {"reported": {}, "suppressed": True}
    assert any("does not validate" in note for note in result["limitations"])
    assert result["model_result_counts"]["suppressed"] is True
    assert "parse_error" not in result["model_result_counts"]["reported"]
    assert result["sample_coverage"]["by_stratum"]["cv_language"] == {"reported": {}, "suppressed": True}
    assert result["reviewer_agreement_by_stratum"]["source_format"] == {"reported": {}, "suppressed": True}
    assert result["model_vs_adjudication_by_stratum"]["job_family"] == {"reported": {}, "suppressed": True}
    assert result["stratum_privacy"]["minimum_unique_candidates"] == 5
    assert any("minimum-cell rule" in note for note in result["limitations"])
    serialized_result = json.dumps(result)
    assert not re.search(r"\b(?:job|cand|ev)_[a-f0-9]{12}\b", serialized_result)
    assert "crit_0123456789ab" in result["model_vs_adjudication_by_stratum"]["criterion_id"]["reported"]
    assert result["model_vs_adjudication_by_stratum"]["criterion_id"]["reported"]["crit_0123456789ab"]["accuracy"] is None
    assert result["model_vs_adjudication_by_stratum"]["criterion_id"]["reported"]["crit_0123456789ab"]["accuracy_suppressed"] is True
    assert "software_engineering" not in serialized_result
    assert "product_design" not in serialized_result
    assert "person@example.test" not in serialized_result
    assert eligible_result["sample_coverage"]["by_stratum"]["job_family"]["reported"]["software_engineering"] == 5
    assert eligible_result["reviewer_agreement_by_stratum"]["job_family"]["reported"]["software_engineering"]["double_reviewed_rows"] == 5
    assert eligible_result["model_vs_adjudication_by_stratum"]["job_family"]["reported"]["software_engineering"]["compared_rows"] == 5
    assert eligible_result["model_vs_adjudication"]["by_evidence_label"]["supported"]["support"] == 5
    assert eligible_result["model_vs_adjudication"]["accuracy"] == 1.0
    assert eligible_result["reviewer_agreement"]["exact_agreement_count"] == 5
    assert eligible_result["evidence_reference_review"]["adjudicated"]["status_counts"]["valid"] == 5
    assert eligible_result["evidence_reference_review"]["adjudicated"]["citation_validity_rate"] == 1.0
    assert incorrect_reference_result["evidence_reference_review"]["adjudicated"]["status_counts"]["incorrect"] == 5
    assert incorrect_reference_result["evidence_reference_review"]["adjudicated"]["citation_validity_rate"] == 0.0
    assert single_sparse_outcome_result["overall_privacy"]["sample_suppressed"] is False
    assert single_sparse_outcome_result["evidence_locator_presence"]["reference"]["suppressed"] is True
    assert single_sparse_outcome_result["evidence_locator_presence"]["page_number"]["suppressed"] is True
    assert single_sparse_outcome_result["evidence_locator_presence"]["reference"]["by_stratum"]["criterion_id"] == {
        "reported": {}, "suppressed": True,
    }
    single_sparse_reference = single_sparse_outcome_result["evidence_reference_review"]["adjudicated"]
    assert single_sparse_reference["status_counts"]["valid"] is None
    assert single_sparse_reference["status_counts"]["missing"] is None
    assert single_sparse_reference["status_counts_suppressed"] is True
    assert single_sparse_reference["citation_coverage_rate"] is None
    assert partial_metrics_result["overall_privacy"]["sample_suppressed"] is False
    assert partial_metrics_result["double_reviewed_rows"] is None
    assert partial_metrics_result["adjudicated_rows"] is None
    assert partial_metrics_result["reviewer_agreement"] is None
    assert partial_metrics_result["model_vs_adjudication"] is None
    assert partial_metrics_result["evidence_reference_review"]["review_coverage"]["reviewer_double_reviewed_rows"] is None
    assert partial_metrics_result["evidence_reference_review"]["review_coverage"]["adjudicated_rows"] is None
    assert partial_metrics_result["overall_privacy"]["suppressed"] is True
    assert "model_vs_adjudication" in partial_metrics_result["overall_privacy"]["suppressed_measures"]
    assert sparse_result["sample_coverage"]["by_stratum"]["job_family"] == {"reported": {}, "suppressed": True}
    assert sparse_result["model_vs_adjudication_by_stratum"]["job_family"] == {"reported": {}, "suppressed": True}
    assert repeated_candidate_result["sample_coverage"]["by_stratum"]["job_family"] == {"reported": {}, "suppressed": True}
    assert repeated_candidate_result["reviewer_agreement_by_stratum"]["job_family"] == {"reported": {}, "suppressed": True}
    assert repeated_candidate_result["model_vs_adjudication_by_stratum"]["job_family"] == {"reported": {}, "suppressed": True}
    assert repeated_candidate_result["model_vs_adjudication_by_stratum"]["criterion_id"] == {"reported": {}, "suppressed": True}
    assert repeated_candidate_result["stratum_privacy"]["candidate_count_unit"].startswith("distinct candidate_id values")
    assert repeated_candidate_result["overall_privacy"]["suppressed"] is True
    assert repeated_candidate_result["overall_privacy"]["distinct_candidate_count"] is None
    assert repeated_candidate_result["input_rows"] is None
    assert repeated_candidate_result["sample_coverage"]["rows"] is None
    assert repeated_candidate_result["double_reviewed_rows"] is None
    assert repeated_candidate_result["adjudicated_rows"] is None
    assert repeated_candidate_result["model_result_counts"] == {"reported": {}, "suppressed": True}
    assert repeated_candidate_result["reviewer_agreement"] is None
    assert repeated_candidate_result["model_vs_adjudication"] is None
    assert repeated_candidate_result["evidence_locator_presence"]["reference"]["suppressed"] is True
    assert repeated_candidate_result["evidence_locator_presence"]["page_number"]["suppressed"] is True
    assert repeated_candidate_result["evidence_reference_review"]["review_coverage"]["reviewer_double_reviewed_rows"] is None
    assert "reviewer_agreement" in repeated_candidate_result["overall_privacy"]["suppressed_measures"]
    assert "software_engineering" not in json.dumps(sparse_result)
    assert "product_design" not in json.dumps(sparse_result)
    try:
        evaluate_adjudications.evaluate(ROOT / "data" / "evaluation" / "annotation-template.csv")
    except ValueError as error:
        assert "outside the repository" in str(error)
    else:
        raise AssertionError("annotation templates inside the repository must be rejected")
    with tempfile.TemporaryDirectory(prefix="karsahire-empty-annotation-") as directory:
        empty_path = Path(directory) / "annotation-template.csv"
        with empty_path.open("w", encoding="utf-8", newline="") as file:
            csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS).writeheader()
        try:
            evaluate_adjudications.evaluate(empty_path)
        except ValueError as error:
            assert "belum memiliki baris anotasi" in str(error)
        else:
            raise AssertionError("header-only external annotation template should not produce evaluation metrics")


def check_blind_annotation_packets() -> None:
    with tempfile.TemporaryDirectory(prefix="karsahire-blind-review-") as directory:
        base = Path(directory)
        master_dir = base / "master"
        recruiter_dir = base / "recruiter-private"
        manager_dir = base / "manager-private"
        merged_dir = base / "evaluation-private"
        for folder in (master_dir, recruiter_dir, manager_dir, merged_dir):
            folder.mkdir()

        master_path = master_dir / "synthetic-master.csv"
        master_rows = []
        for index in (1, 2):
            row = {field: "" for field in evaluate_adjudications.FIELDS}
            row.update({
                "requisition_id": "job_000000000001",
                "candidate_id": f"cand_{index:012x}",
                "criteria_version": "v1",
                "criterion_id": "crit_000000000001",
                "pipeline_version": "local-development",
                "job_family": "software_engineering",
                "cv_language": "en",
                "source_format": "pdf_text",
                "model_result": "matched",
                "model_evidence_ref": f"ev_{index:012x}",
                "model_evidence_page": "1",
            })
            master_rows.append(row)
        with master_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            writer.writerows(master_rows)

        preparation_binding = master_dir / "preparation-binding.json"
        run_id = annotation_packets.freeze_master(master_path, preparation_binding)
        recruiter_path = recruiter_dir / "review.csv"
        manager_path = manager_dir / "review.csv"
        assert annotation_packets.prepare_packet(
            master_path, recruiter_path, "recruiter", preparation_binding
        ) == 2
        assert annotation_packets.prepare_packet(
            master_path, manager_path, "hiring_manager", preparation_binding
        ) == 2

        changed_master = master_dir / "changed-master.csv"
        changed_rows = [dict(row) for row in master_rows]
        changed_rows[0]["model_evidence_ref"] = "ev_0000000000ff"
        with changed_master.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            writer.writerows(changed_rows)
        try:
            annotation_packets.prepare_packet(
                changed_master, recruiter_dir / "changed.csv", "recruiter", preparation_binding
            )
        except ValueError as error:
            assert "changed after it was frozen" in str(error)
        else:
            raise AssertionError("annotation preparation must reject changed model evidence after freezing")
        for packet_path, expected_role, labels_by_candidate in (
            (recruiter_path, "recruiter", {
                "cand_000000000001": ("supported", "valid"),
                "cand_000000000002": ("partial", "valid"),
            }),
            (manager_path, "hiring_manager", {
                "cand_000000000001": ("partial", "incorrect"),
                "cand_000000000002": ("supported", "valid"),
            }),
        ):
            with packet_path.open("r", encoding="utf-8", newline="") as file:
                reader = csv.DictReader(file)
                assert tuple(reader.fieldnames or ()) == annotation_packets.PACKET_FIELDS
                packet_rows = list(reader)
            assert "model_result" not in packet_rows[0]
            assert "model_evidence_ref" not in packet_rows[0]
            assert "model_evidence_page" not in packet_rows[0]
            assert "adjudicated_label" not in packet_rows[0]
            assert all(row["reviewer_role"] == expected_role for row in packet_rows)
            assert {row["run_id"] for row in packet_rows} == {run_id}
            for row in packet_rows:
                label, reference_label = labels_by_candidate[row["candidate_id"]]
                row["reviewer_label"] = label
                row["reference_label"] = reference_label
            with packet_path.open("w", encoding="utf-8", newline="") as file:
                writer = csv.DictWriter(file, fieldnames=annotation_packets.PACKET_FIELDS)
                writer.writeheader()
                writer.writerows(packet_rows)

        second_binding = master_dir / "second-preparation-binding.json"
        annotation_packets.freeze_master(master_path, second_binding)
        mixed_manager_path = manager_dir / "other-run.csv"
        assert annotation_packets.prepare_packet(
            master_path, mixed_manager_path, "hiring_manager", second_binding
        ) == 2
        try:
            annotation_packets.merge_packets(
                master_path, recruiter_path, mixed_manager_path,
                merged_dir / "mixed.csv", preparation_binding,
            )
        except ValueError as error:
            assert "run ID does not match" in str(error)
        else:
            raise AssertionError("annotation merger must reject packets from different runs")

        merged_path = merged_dir / "merged.csv"
        assert annotation_packets.merge_packets(
            master_path, recruiter_path, manager_path, merged_path, preparation_binding
        ) == 2
        with merged_path.open("r", encoding="utf-8", newline="") as file:
            merged_rows = list(csv.DictReader(file))
        assert merged_rows[0]["recruiter_label"] == "supported"
        assert merged_rows[0]["hiring_manager_label"] == "partial"
        assert merged_rows[0]["recruiter_reference_label"] == "valid"
        assert merged_rows[0]["hiring_manager_reference_label"] == "incorrect"
        assert merged_rows[0]["adjudicated_label"] == ""
        merged_result = evaluate_adjudications.evaluate(merged_path)
        assert merged_result["double_reviewed_rows"] is None
        assert merged_result["overall_privacy"]["sample_suppressed"] is True

        try:
            annotation_packets.merge_packets(
                master_path, recruiter_path, recruiter_path, merged_dir / "duplicate-role.csv",
                preparation_binding,
            )
        except ValueError as error:
            assert "separate packet" in str(error)
        else:
            raise AssertionError("packet merger must require distinct reviewer packet files")

        truncated_packet = manager_dir / "truncated.csv"
        with truncated_packet.open("w", encoding="utf-8", newline="") as file:
            writer = csv.writer(file)
            writer.writerow(annotation_packets.PACKET_FIELDS)
            writer.writerow(("incomplete",))
        try:
            annotation_packets.packet_map(truncated_packet, "hiring_manager", run_id)
        except ValueError as error:
            assert "missing or extra columns" in str(error)
        else:
            raise AssertionError("packet reader must reject rows with missing fields")

        labeled_master = master_dir / "labeled-master.csv"
        master_rows[0]["recruiter_label"] = "supported"
        with labeled_master.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_adjudications.FIELDS)
            writer.writeheader()
            writer.writerows(master_rows)
        try:
            annotation_packets.prepare_packet(
                labeled_master, recruiter_dir / "must-not-exist.csv", "recruiter", preparation_binding
            )
        except ValueError as error:
            assert "pristine" in str(error)
        else:
            raise AssertionError("blind packet preparation must reject an already-labeled master")

        try:
            annotation_packets.require_outside_repository(ROOT / "data" / "evaluation" / "annotation-template.csv")
        except ValueError as error:
            assert "outside the repository" in str(error)
        else:
            raise AssertionError("annotation data must not be written inside the repository")
    print("PASS: blind annotation packets freeze their master, randomize rows, and reject mixed runs")


def check_blind_ranking_packets() -> None:
    with tempfile.TemporaryDirectory(prefix="karsahire-ranking-review-") as directory:
        base = Path(directory)
        operator_dir = base / "operator"
        recruiter_dir = base / "recruiter"
        manager_dir = base / "manager"
        adjudicator_dir = base / "adjudicator"
        for folder in (operator_dir, recruiter_dir, manager_dir, adjudicator_dir):
            folder.mkdir()

        master_path = operator_dir / "master.csv"
        master_rows = []
        for index in range(1, 6):
            row = {field: "" for field in evaluate_rankings.FIELDS}
            row.update({
                "requisition_id": "job_000000000001",
                "candidate_id": f"cand_{index:012x}",
                "criteria_version": "v1",
                "pipeline_version": "local-synthetic-test",
                "job_family": "software_engineering",
                "cv_language": "en",
                "source_format": "pdf_text",
                "model_rank": str(index),
            })
            master_rows.append(row)
        with master_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_rankings.FIELDS)
            writer.writeheader()
            writer.writerows(master_rows)

        preparation_binding = operator_dir / "preparation-binding.json"
        run_id = ranking_packets.freeze_master(master_path, preparation_binding)
        assert re.fullmatch(r"[a-f0-9]{32}", run_id)
        recruiter_path = recruiter_dir / "packet.csv"
        manager_path = manager_dir / "packet.csv"
        ranking_packets.prepare(master_path, "recruiter", recruiter_path, preparation_binding)
        ranking_packets.prepare(master_path, "hiring_manager", manager_path, preparation_binding)

        changed_master = operator_dir / "changed-master.csv"
        changed_rows = [dict(row) for row in master_rows]
        changed_rows[0]["model_rank"], changed_rows[-1]["model_rank"] = "5", "1"
        with changed_master.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_rankings.FIELDS)
            writer.writeheader()
            writer.writerows(changed_rows)
        try:
            ranking_packets.prepare(
                changed_master, "recruiter", recruiter_dir / "changed.csv", preparation_binding
            )
        except ValueError as error:
            assert "changed after its order was frozen" in str(error)
        else:
            raise AssertionError("ranking packet preparation must reject a changed frozen model order")

        reviewer_ranks = {
            "recruiter": {f"cand_{index:012x}": index for index in range(1, 6)},
            "hiring_manager": {
                "cand_000000000001": 1,
                "cand_000000000002": 1,
                "cand_000000000003": 3,
                "cand_000000000004": 4,
                "cand_000000000005": 5,
            },
        }
        for packet_path, reviewer_role in (
            (recruiter_path, "recruiter"),
            (manager_path, "hiring_manager"),
        ):
            with packet_path.open("r", encoding="utf-8", newline="") as file:
                reader = csv.DictReader(file)
                assert tuple(reader.fieldnames or ()) == ranking_packets.PACKET_FIELDS
                packet_rows = list(reader)
            assert "model_rank" not in packet_rows[0] and "pipeline_version" not in packet_rows[0]
            assert {row["run_id"] for row in packet_rows} == {run_id}
            for row in packet_rows:
                row["reviewer_rank"] = str(reviewer_ranks[reviewer_role][row["candidate_id"]])
            with packet_path.open("w", encoding="utf-8", newline="") as file:
                writer = csv.DictWriter(file, fieldnames=ranking_packets.PACKET_FIELDS)
                writer.writeheader()
                writer.writerows(packet_rows)

        second_binding = operator_dir / "second-preparation-binding.json"
        ranking_packets.freeze_master(master_path, second_binding)
        mixed_manager_path = manager_dir / "other-run.csv"
        ranking_packets.prepare(master_path, "hiring_manager", mixed_manager_path, second_binding)
        with mixed_manager_path.open("r", encoding="utf-8", newline="") as file:
            reader = csv.DictReader(file)
            mixed_rows = list(reader)
        for row in mixed_rows:
            row["reviewer_rank"] = str(reviewer_ranks["hiring_manager"][row["candidate_id"]])
        with mixed_manager_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=ranking_packets.PACKET_FIELDS)
            writer.writeheader()
            writer.writerows(mixed_rows)
        try:
            ranking_packets.merge(
                master_path, recruiter_path, mixed_manager_path, preparation_binding,
                adjudicator_dir / "mixed.csv", operator_dir / "mixed-binding.json",
            )
        except ValueError as error:
            assert "run ID mismatch" in str(error)
        else:
            raise AssertionError("ranking packet merger must reject packets from different runs")

        adjudication_path = adjudicator_dir / "adjudication.csv"
        final_binding = operator_dir / "master-binding.json"
        ranking_packets.merge(
            master_path, recruiter_path, manager_path, preparation_binding,
            adjudication_path, final_binding,
        )
        adjudicated_by_candidate = {f"cand_{index:012x}": index for index in range(1, 6)}
        with adjudication_path.open("r", encoding="utf-8", newline="") as file:
            reader = csv.DictReader(file)
            adjudication_rows = list(reader)
        assert tuple(reader.fieldnames or ()) == ranking_packets.ADJUDICATION_FIELDS
        assert all(row["run_id"] == run_id for row in adjudication_rows)
        assert "model_rank" not in adjudication_rows[0] and "pipeline_version" not in adjudication_rows[0]
        for row in adjudication_rows:
            row["adjudicated_rank"] = str(adjudicated_by_candidate[row["candidate_id"]])
        with adjudication_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=ranking_packets.ADJUDICATION_FIELDS)
            writer.writeheader()
            writer.writerows(adjudication_rows)

        evaluation_path = operator_dir / "evaluation.csv"
        ranking_packets.apply_adjudication(
            master_path, adjudication_path, final_binding, evaluation_path
        )
        report = evaluate_rankings.evaluate(evaluation_path)
        assert report["overall"]["model_vs_adjudicated_kendall_tau_b"]["macro_mean"] == 1.0
        assert report["overall"]["model_top_k_overlap_precision"]["macro_mean"] == 1.0
        assert report["overall"]["model_top_k_overlap_recall"]["macro_mean"] == 1.0

        suppression_path = operator_dir / "ranking-suppression.csv"
        suppression_rows = []
        for index in range(1, 13):
            row = {field: "" for field in evaluate_rankings.FIELDS}
            row.update({
                "requisition_id": "job_000000000002",
                "candidate_id": f"cand_{index:012x}",
                "criteria_version": "v1",
                "pipeline_version": "local-synthetic-test",
                "job_family": "software_engineering",
                "cv_language": "en" if index <= 5 else "id" if index <= 10 else "fr",
                "source_format": "pdf_text" if index <= 5 else "docx" if index <= 10 else "image",
                "model_rank": str(index),
                "recruiter_rank": str(index),
                "hiring_manager_rank": str(index),
                "adjudicated_rank": str(index),
            })
            suppression_rows.append(row)
        with suppression_path.open("w", encoding="utf-8", newline="") as file:
            writer = csv.DictWriter(file, fieldnames=evaluate_rankings.FIELDS)
            writer.writeheader()
            writer.writerows(suppression_rows)
        suppressed_report = evaluate_rankings.evaluate(suppression_path)
        language_slice = suppressed_report["by_stratum"]["cv_language"]
        format_slice = suppressed_report["by_stratum"]["source_format"]
        assert language_slice["some_metrics_suppressed"] is True
        assert set(language_slice["reported"]) == {"en"}
        assert format_slice["some_metrics_suppressed"] is True
        assert set(format_slice["reported"]) == {"docx"}
        assert suppressed_report["by_stratum"]["job_family"]["some_metrics_suppressed"] is False
    print("PASS: ranking packets freeze the model order, bind both reviewers to one run, preserve blind adjudication, and complementarily suppress sparse strata")


def check_local_bind_config() -> None:
    with patch.dict(os.environ, {
        "RECRUITMENT_COPILOT_HOST": "127.0.0.1",
        "RECRUITMENT_COPILOT_PORT": "8765",
    }, clear=False):
        assert server.local_bind_config() == ("127.0.0.1", 8765)

    with patch.dict(os.environ, {"RECRUITMENT_COPILOT_HOST": "0.0.0.0"}, clear=False):
        try:
            server.local_bind_config()
        except SystemExit as error:
            assert "hanya boleh bind ke 127.0.0.1" in str(error)
        else:
            raise AssertionError("server must reject a non-loopback bind address without authentication")

    for value in ("not-a-port", "0", "65536"):
        with patch.dict(os.environ, {"RECRUITMENT_COPILOT_PORT": value}, clear=False):
            try:
                server.local_bind_config()
            except SystemExit as error:
                assert "1–65535" in str(error)
            else:
                raise AssertionError("server must reject malformed or out-of-range ports")


def check_shutdown_drains_active_handler() -> None:
    handler_started = threading.Event()
    release_handler = threading.Event()
    client_result = []

    class BlockingHandler(BaseHTTPRequestHandler):
        def do_GET(self) -> None:  # noqa: N802
            if self.path == "/probe":
                self.send_response(200)
                self.send_header("Content-Length", "5")
                self.end_headers()
                self.wfile.write(b"probe")
                return
            handler_started.set()
            if not release_handler.wait(timeout=5):
                self.send_error(500)
                return
            self.send_response(200)
            self.send_header("Content-Length", "7")
            self.end_headers()
            self.wfile.write(b"drained")

        def log_message(self, _format: str, *_args) -> None:
            pass

    httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), BlockingHandler)
    httpd.daemon_threads = False
    httpd.block_on_close = True
    server_thread = threading.Thread(target=httpd.serve_forever, daemon=True)
    server_thread.start()
    server.DRAINING.clear()

    def request() -> None:
        try:
            with urllib.request.urlopen(
                f"http://127.0.0.1:{httpd.server_address[1]}/", timeout=5
            ) as response:
                client_result.append((response.status, response.read()))
        except Exception as error:  # preserve worker failures for the assertion below
            client_result.append(error)

    client_thread = threading.Thread(target=request, daemon=True)
    client_thread.start()
    close_thread = None
    server_closed = False
    try:
        assert handler_started.wait(timeout=2), "request handler did not start"
        signal_log = io.StringIO()
        with redirect_stdout(signal_log):
            assert server.request_shutdown_for_signal(httpd, grace_seconds=1, signum=signal.SIGTERM)
        assert "SIGTERM diterima" in signal_log.getvalue()
        assert server.DRAINING.is_set(), "shutdown did not mark the process as draining"
        with urllib.request.urlopen(
            f"http://127.0.0.1:{httpd.server_address[1]}/probe", timeout=2
        ) as probe:
            assert probe.status == 200 and probe.read() == b"probe"
        server_thread.join(timeout=2)
        assert not server_thread.is_alive(), "serve_forever did not stop"

        close_thread = threading.Thread(target=httpd.server_close)
        close_thread.start()
        close_thread.join(timeout=0.1)
        assert close_thread.is_alive(), "server_close did not wait for its active handler"

        release_handler.set()
        close_thread.join(timeout=2)
        server_closed = True
        client_thread.join(timeout=2)
        assert not close_thread.is_alive(), "server_close stayed blocked after handler exit"
        assert not client_thread.is_alive(), "synthetic request did not finish"
        assert client_result == [(200, b"drained")], f"unexpected handler result: {client_result!r}"
        print("PASS: SIGTERM callback marks drain and waits for an active synthetic request handler")
    finally:
        release_handler.set()
        if server_thread.is_alive():
            httpd.shutdown()
            server_thread.join(timeout=2)
        if close_thread is not None:
            close_thread.join(timeout=2)
        if not server_closed:
            httpd.server_close()
        client_thread.join(timeout=2)
        server.DRAINING.clear()


def check_database_path_config() -> None:
    assert server.database_path_from_env("") == server.DATA / "copilot.sqlite3"
    with tempfile.TemporaryDirectory(prefix="karsahire-db-path-") as directory:
        custom_path = Path(directory) / "private" / "copilot.sqlite3"
        with patch.dict(os.environ, {server.DB_PATH_ENV: str(custom_path)}, clear=False):
            assert server.database_path_from_env() == custom_path
        assert server.database_path_from_env(str(custom_path)) == custom_path

        previous_db = server.DB_PATH
        try:
            server.DB_PATH = custom_path
            with server.connect() as db:
                assert db.execute("SELECT 1").fetchone()[0] == 1
            assert custom_path.is_file()
        finally:
            server.DB_PATH = previous_db

    try:
        server.database_path_from_env("relative/copilot.sqlite3")
    except SystemExit as error:
        assert server.DB_PATH_ENV in str(error)
    else:
        raise AssertionError("database path override must require an absolute path")


def check_shutdown_grace_config() -> None:
    with patch.dict(os.environ, {}, clear=False):
        os.environ.pop(server.SHUTDOWN_GRACE_ENV, None)
        assert server.shutdown_grace_config() == server.DEFAULT_SHUTDOWN_GRACE_SECONDS
    for value, expected in (("0", 0), ("12", 12), ("120", 120)):
        with patch.dict(os.environ, {server.SHUTDOWN_GRACE_ENV: value}, clear=False):
            assert server.shutdown_grace_config() == expected
    for value in ("-1", "121", "invalid"):
        with patch.dict(os.environ, {server.SHUTDOWN_GRACE_ENV: value}, clear=False):
            try:
                server.shutdown_grace_config()
            except SystemExit as error:
                assert server.SHUTDOWN_GRACE_ENV in str(error)
            else:
                raise AssertionError("shutdown grace must reject invalid or out-of-range values")


def check_readiness_during_drain(base: str) -> None:
    server.DRAINING.set()
    try:
        status, live = request_json(f"{base}/api/health/live")
        assert status == 200 and live["live"] and live["draining"]

        try:
            request_json(f"{base}/api/health/ready")
        except urllib.error.HTTPError as error:
            ready = json.loads(error.read().decode("utf-8"))
            assert error.code == 503 and ready["ready"] is False
            assert ready["draining"] is True and ready["database"] == "ready"
        else:
            raise AssertionError("readiness must fail while the service is draining")

        try:
            request_json(f"{base}/api/jobs")
        except urllib.error.HTTPError as error:
            assert error.code == 503
            error.close()
        else:
            raise AssertionError("normal reads must stop while the service is draining")

        try:
            request_json(
                f"{base}/api/jobs",
                "POST",
                {"title": "Must not be created while draining", "criteria": []},
            )
        except urllib.error.HTTPError as error:
            assert error.code == 503
            error.close()
        else:
            raise AssertionError("mutations must stop while the service is draining")
    finally:
        server.DRAINING.clear()
    status, jobs = request_json(f"{base}/api/jobs")
    assert status == 200 and jobs == []
    print("PASS: draining keeps liveness, fails readiness, and rejects new API work")


def check_schema_migration_and_database_backup() -> None:
    previous_data, previous_db = server.DATA, server.DB_PATH
    try:
        with tempfile.TemporaryDirectory(prefix="karsahire-db-ops-") as directory:
            root = Path(directory)
            server.DATA = root
            source = root / "legacy.sqlite3"
            server.DB_PATH = source
            db = sqlite3.connect(source)
            try:
                db.execute(
                    """CREATE TABLE jobs (
                        id TEXT PRIMARY KEY,
                        title TEXT NOT NULL,
                        department TEXT NOT NULL DEFAULT '',
                        description TEXT NOT NULL DEFAULT '',
                        criteria_json TEXT NOT NULL,
                        created_at TEXT NOT NULL
                    )"""
                )
                db.execute(
                    "INSERT INTO jobs(id,title,department,description,criteria_json,created_at) "
                    "VALUES(?,?,?,?,?,?)",
                    ("job_synthetic", "Synthetic migration fixture", "", "", "[]", "2026-01-01T00:00:00+00:00"),
                )
                db.commit()
            finally:
                db.close()

            server.init_db()
            with server.connect() as migrated:
                assert migrated.execute("PRAGMA user_version").fetchone()[0] == server.SCHEMA_VERSION
                row = migrated.execute("SELECT title, status FROM jobs WHERE id=?", ("job_synthetic",)).fetchone()
                assert row["title"] == "Synthetic migration fixture"
                assert row["status"] == "awaiting_approval"

            backup_path = root / "backups" / "snapshot.sqlite3"
            result = backup_database.create_backup(source, backup_path)
            assert result["integrity_check"] == "ok"
            assert result["schema_version"] == server.SCHEMA_VERSION

            repository_backup = ROOT / "data" / "smoke-rejected" / "snapshot.sqlite3"
            try:
                backup_database.create_backup(source, repository_backup)
            except ValueError as error:
                assert "di luar repository" in str(error)
            else:
                raise AssertionError("backup must refuse a destination inside the repository")
            assert not repository_backup.parent.exists(), "unsafe backup must be rejected before creating folders"

            backup_conn = sqlite3.connect(backup_path)
            try:
                assert backup_conn.execute("PRAGMA integrity_check").fetchone() == ("ok",)
                assert backup_conn.execute("SELECT COUNT(*) FROM jobs").fetchone() == (1,)
            finally:
                backup_conn.close()

            try:
                backup_database.create_backup(source, backup_path)
            except FileExistsError:
                pass
            else:
                raise AssertionError("backup command must refuse to overwrite an existing snapshot")

            restored = root / "restored.sqlite3"
            restore_result = restore_database.restore_snapshot(backup_path, restored)
            assert restore_result["integrity_check"] == "ok"

            repository_restore = ROOT / "data" / "smoke-rejected" / "restore.sqlite3"
            try:
                restore_database.restore_snapshot(backup_path, repository_restore)
            except ValueError as error:
                assert "di luar repository" in str(error)
            else:
                raise AssertionError("restore must refuse a destination inside the repository")
            assert not repository_restore.parent.exists(), "unsafe restore must be rejected before creating folders"

            server.DB_PATH = restored
            server.init_db()
            with server.connect() as restored_db:
                assert restored_db.execute("PRAGMA integrity_check").fetchone()[0] == "ok"
                assert restored_db.execute("SELECT title FROM jobs WHERE id=?", ("job_synthetic",)).fetchone()[0] == "Synthetic migration fixture"

            future_snapshot = root / "future-schema.sqlite3"
            future_db = sqlite3.connect(future_snapshot)
            try:
                future_db.execute(f"PRAGMA user_version = {server.SCHEMA_VERSION + 1}")
                future_db.commit()
            finally:
                future_db.close()
            unsupported_restore = root / "unsupported-restore.sqlite3"
            try:
                restore_database.restore_snapshot(future_snapshot, unsupported_restore)
            except ValueError as error:
                assert "lebih baru" in str(error)
            else:
                raise AssertionError("restore must reject schema versions newer than the app supports")
            assert not unsupported_restore.exists()
    finally:
        server.DATA, server.DB_PATH = previous_data, previous_db


def check_local_restore_cutover_and_rollback() -> None:
    """Exercise the documented local database-path switch with synthetic jobs."""
    previous_data, previous_db = server.DATA, server.DB_PATH
    previous_uploads_enabled = server.MANUAL_UPLOADS_ENABLED
    active_server = None
    active_thread = None

    def start_service():
        httpd = server.BoundedThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
        httpd.daemon_threads = False
        httpd.block_on_close = True
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{httpd.server_address[1]}"
        return httpd, thread, base

    def stop_service(httpd, thread) -> None:
        httpd.shutdown()
        thread.join(timeout=3)
        assert not thread.is_alive(), "local cutover service did not stop"
        httpd.server_close()

    try:
        with tempfile.TemporaryDirectory(prefix="karsahire-cutover-") as directory:
            root = Path(directory)
            primary = root / "primary" / "copilot.sqlite3"
            snapshot = root / "backups" / "primary-snapshot.sqlite3"
            restored = root / "staging" / "copilot.sqlite3"
            server.DATA = primary.parent
            server.DB_PATH = primary
            server.MANUAL_UPLOADS_ENABLED = False
            server.DRAINING.clear()
            server.init_db()

            active_server, active_thread, base = start_service()
            status, readiness = request_json(f"{base}/api/health/ready")
            assert status == 200 and readiness["ready"] is True
            status, original_job = request_json(
                f"{base}/api/jobs",
                "POST",
                {
                    "title": "Synthetic cutover source",
                    "department": "Local QA",
                    "description": "Original record for local restore rehearsal.",
                    "criteria": [{"type": "required", "label": "Python"}],
                },
            )
            assert status == 201
            original_id = original_job["id"]
            backup_database.create_backup(primary, snapshot)
            restore_database.restore_snapshot(snapshot, restored)

            stop_service(active_server, active_thread)
            active_server = active_thread = None
            server.DB_PATH = restored
            server.DATA = restored.parent
            server.init_db()
            active_server, active_thread, base = start_service()
            status, readiness = request_json(f"{base}/api/health/ready")
            assert status == 200 and readiness["ready"] is True
            status, restored_jobs = request_json(f"{base}/api/jobs")
            assert status == 200 and [job["id"] for job in restored_jobs] == [original_id]

            status, stage_job = request_json(
                f"{base}/api/jobs",
                "POST",
                {
                    "title": "Synthetic cutover-only record",
                    "department": "Local QA",
                    "description": "Must disappear when the local switch is rolled back.",
                    "criteria": [{"type": "required", "label": "Python"}],
                },
            )
            assert status == 201
            stage_id = stage_job["id"]
            stop_service(active_server, active_thread)
            active_server = active_thread = None

            server.DB_PATH = primary
            server.DATA = primary.parent
            server.init_db()
            active_server, active_thread, base = start_service()
            status, readiness = request_json(f"{base}/api/health/ready")
            assert status == 200 and readiness["ready"] is True
            status, rolled_back_jobs = request_json(f"{base}/api/jobs")
            assert status == 200
            rolled_back_ids = {job["id"] for job in rolled_back_jobs}
            assert rolled_back_ids == {original_id} and stage_id not in rolled_back_ids
            print("PASS: local synthetic snapshot restore, database-path cutover, readiness, and rollback")
    finally:
        if active_server is not None and active_thread is not None:
            stop_service(active_server, active_thread)
        server.DRAINING.clear()
        server.DATA, server.DB_PATH = previous_data, previous_db
        server.MANUAL_UPLOADS_ENABLED = previous_uploads_enabled


def check_process_restore_cutover_and_rollback() -> None:
    """Reopen the configured database path across real local service processes."""
    child_code = "\n".join((
        "import os, signal, threading, time",
        "from pathlib import Path",
        f"import sys; sys.path.insert(0, {str(ROOT)!r})",
        "import server",
        "server.init_db()",
        "host, port = server.local_bind_config()",
        "grace = server.shutdown_grace_config()",
        "httpd = server.BoundedThreadingHTTPServer((host, port), server.Handler)",
        "httpd.daemon_threads = False",
        "httpd.block_on_close = True",
        "server.register_shutdown_signal_handlers(httpd, grace)",
        "stop_file = Path(os.environ['KARSAHIRE_SMOKE_STOP_FILE'])",
        "def stop_when_requested():",
        "    while not stop_file.exists(): time.sleep(0.02)",
        "    server.request_shutdown_for_signal(httpd, grace, signal.SIGTERM)",
        "threading.Thread(target=stop_when_requested, daemon=True).start()",
        "try: httpd.serve_forever()",
        "finally:",
        "    server.DRAINING.set()",
        "    httpd.server_close()",
    ))

    def new_port() -> int:
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
            listener.bind(("127.0.0.1", 0))
            return listener.getsockname()[1]

    def start_process(database: Path, stop_file: Path):
        port = new_port()
        child_env = local_release.child_environment()
        child_env.update({
            server.DB_PATH_ENV: str(database),
            "RECRUITMENT_COPILOT_HOST": "127.0.0.1",
            "RECRUITMENT_COPILOT_PORT": str(port),
            "RECRUITMENT_COPILOT_SHUTDOWN_GRACE_SECONDS": "0",
            server.MANUAL_UPLOADS_ENV: "false",
            "KARSAHIRE_SMOKE_STOP_FILE": str(stop_file),
        })
        process = subprocess.Popen(
            [sys.executable, "-c", child_code],
            cwd=ROOT,
            env=child_env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
        )
        base = f"http://127.0.0.1:{port}"
        deadline = time.monotonic() + 10
        while time.monotonic() < deadline:
            if process.poll() is not None:
                stdout, stderr = process.communicate()
                raise AssertionError(
                    f"isolated service process exited during startup ({process.returncode}): "
                    f"{(stdout + stderr)[-1200:]}"
                )
            try:
                status, readiness = request_json(f"{base}/api/health/ready")
                if status == 200 and readiness["ready"] is True:
                    return process, base
            except (OSError, urllib.error.URLError):
                pass
            time.sleep(0.05)
        stop_file.touch()
        try:
            stdout, stderr = process.communicate(timeout=5)
        except subprocess.TimeoutExpired:
            process.kill()
            stdout, stderr = process.communicate()
        raise AssertionError(
            f"isolated service process did not become ready: {(stdout + stderr)[-1200:]}"
        )

    def stop_process(process, stop_file: Path) -> None:
        stop_file.touch()
        try:
            stdout, stderr = process.communicate(timeout=8)
        except subprocess.TimeoutExpired as error:
            process.kill()
            process.communicate()
            raise AssertionError("isolated service process did not finish its graceful drain") from error
        assert process.returncode == 0, f"isolated service process failed to stop: {stderr[-1200:]}"
        assert "SIGTERM diterima" in stdout, "synthetic stop callback did not run in the child process"

    def stop_if_running(process, stop_file: Path) -> None:
        if process.poll() is None:
            stop_file.touch()
            try:
                process.communicate(timeout=8)
            except subprocess.TimeoutExpired:
                process.kill()
                process.communicate()

    with tempfile.TemporaryDirectory(prefix="karsahire-process-cutover-") as directory:
        root = Path(directory)
        primary = root / "primary" / "copilot.sqlite3"
        snapshot = root / "backups" / "primary-snapshot.sqlite3"
        restored = root / "restored" / "copilot.sqlite3"
        process = None
        stop_file = root / "stop-primary-before-backup"
        try:
            process, base = start_process(primary, stop_file)
            status, original_job = request_json(
                f"{base}/api/jobs",
                "POST",
                {
                    "title": "Synthetic process cutover source",
                    "department": "Local QA",
                    "description": "Record used to verify a process-level restore rehearsal.",
                    "criteria": [{"type": "required", "label": "Python"}],
                },
            )
            assert status == 201
            original_id = original_job["id"]
            backup_database.create_backup(primary, snapshot)
            restore_database.restore_snapshot(snapshot, restored)
            stop_process(process, stop_file)
            process = None

            stop_file = root / "stop-restored-process"
            process, base = start_process(restored, stop_file)
            status, restored_jobs = request_json(f"{base}/api/jobs")
            assert status == 200 and [job["id"] for job in restored_jobs] == [original_id]
            status, stage_job = request_json(
                f"{base}/api/jobs",
                "POST",
                {
                    "title": "Synthetic process cutover only",
                    "department": "Local QA",
                    "description": "This record is discarded by restoring the original path.",
                    "criteria": [{"type": "required", "label": "Python"}],
                },
            )
            assert status == 201
            stage_id = stage_job["id"]
            stop_process(process, stop_file)
            process = None

            stop_file = root / "stop-rolled-back-process"
            process, base = start_process(primary, stop_file)
            status, rolled_back_jobs = request_json(f"{base}/api/jobs")
            assert status == 200
            rolled_back_ids = {job["id"] for job in rolled_back_jobs}
            assert rolled_back_ids == {original_id} and stage_id not in rolled_back_ids
            stop_process(process, stop_file)
            process = None
            print("PASS: separate synthetic service processes reopen the configured DB path after restore and rollback")
        finally:
            if process is not None:
                stop_if_running(process, stop_file)


def check_windows_ctrl_c_process_shutdown() -> None:
    """Deliver a real Ctrl+C console event to a disposable Windows service process."""
    if os.name != "nt":
        print("SKIP: Windows Ctrl+C process delivery is only available on Windows")
        return

    child_code = "\n".join((
        "import ctypes, sys",
        "if not ctypes.windll.kernel32.SetConsoleCtrlHandler(None, False):",
        "    raise OSError(ctypes.get_last_error(), 'could not enable Ctrl+C handling')",
        f"sys.path.insert(0, {str(ROOT)!r})",
        "import server",
        "server.main()",
    ))
    sender_code = "\n".join((
        "import ctypes, sys",
        "from ctypes import wintypes",
        "kernel = ctypes.WinDLL('kernel32', use_last_error=True)",
        "kernel.AttachConsole.argtypes = [wintypes.DWORD]",
        "kernel.AttachConsole.restype = wintypes.BOOL",
        "kernel.SetConsoleCtrlHandler.argtypes = [wintypes.HANDLE, wintypes.BOOL]",
        "kernel.SetConsoleCtrlHandler.restype = wintypes.BOOL",
        "kernel.GenerateConsoleCtrlEvent.argtypes = [wintypes.DWORD, wintypes.DWORD]",
        "kernel.GenerateConsoleCtrlEvent.restype = wintypes.BOOL",
        "kernel.FreeConsole()",
        "if not kernel.AttachConsole(int(sys.argv[1])):",
        "    raise OSError(ctypes.get_last_error(), 'could not attach to isolated service console')",
        "try:",
        "    if not kernel.SetConsoleCtrlHandler(None, True):",
        "        raise OSError(ctypes.get_last_error(), 'could not ignore Ctrl+C in event sender')",
        "    if not kernel.GenerateConsoleCtrlEvent(0, 0):",
        "        raise OSError(ctypes.get_last_error(), 'could not send Ctrl+C console event')",
        "finally:",
        "    kernel.FreeConsole()",
    ))

    with tempfile.TemporaryDirectory(prefix="karsahire-ctrl-c-") as directory:
        root = Path(directory)
        database = root / "copilot.sqlite3"
        port = None
        with socket.socket(socket.AF_INET, socket.SOCK_STREAM) as listener:
            listener.bind(("127.0.0.1", 0))
            port = listener.getsockname()[1]
        child_env = local_release.child_environment()
        child_env.update({
            server.DB_PATH_ENV: str(database),
            "RECRUITMENT_COPILOT_HOST": "127.0.0.1",
            "RECRUITMENT_COPILOT_PORT": str(port),
            "RECRUITMENT_COPILOT_SHUTDOWN_GRACE_SECONDS": "0",
            server.MANUAL_UPLOADS_ENV: "false",
        })
        startup = subprocess.STARTUPINFO()
        startup.dwFlags |= subprocess.STARTF_USESHOWWINDOW
        startup.wShowWindow = subprocess.SW_HIDE
        process = subprocess.Popen(
            [sys.executable, "-c", child_code],
            cwd=ROOT,
            env=child_env,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            text=True,
            creationflags=subprocess.CREATE_NEW_CONSOLE | subprocess.CREATE_NEW_PROCESS_GROUP,
            startupinfo=startup,
        )
        try:
            base = f"http://127.0.0.1:{port}"
            deadline = time.monotonic() + 10
            while time.monotonic() < deadline:
                if process.poll() is not None:
                    stdout, stderr = process.communicate()
                    raise AssertionError(
                        f"isolated Ctrl+C service exited during startup ({process.returncode}): "
                        f"{(stdout + stderr)[-1200:]}"
                    )
                try:
                    status, readiness = request_json(f"{base}/api/health/ready")
                    if status == 200 and readiness["ready"] is True:
                        break
                except (OSError, urllib.error.URLError):
                    time.sleep(0.05)
            else:
                raise AssertionError("isolated Ctrl+C service did not become ready")

            status, job = request_json(
                f"{base}/api/jobs",
                "POST",
                {
                    "title": "Synthetic Ctrl+C persistence check",
                    "department": "Local QA",
                    "description": "Temporary lowongan used only to verify graceful local stop.",
                    "criteria": [{"type": "required", "label": "Python"}],
                },
            )
            assert status == 201 and job["id"]
            sender = subprocess.run(
                [sys.executable, "-c", sender_code, str(process.pid)],
                cwd=ROOT,
                env=child_env,
                capture_output=True,
                text=True,
                timeout=5,
                creationflags=subprocess.CREATE_NO_WINDOW,
            )
            assert sender.returncode == 0, f"Ctrl+C sender failed: {sender.stderr[-1000:]}"
            stdout, stderr = process.communicate(timeout=10)
            assert process.returncode == 0, f"service did not exit cleanly: {stderr[-1200:]}"
            assert "SIGINT diterima" in stdout, "service did not run its SIGINT drain callback"

            database_check = sqlite3.connect(database)
            try:
                assert database_check.execute("PRAGMA integrity_check").fetchone() == ("ok",)
                assert database_check.execute("PRAGMA user_version").fetchone() == (server.SCHEMA_VERSION,)
                assert database_check.execute("SELECT COUNT(*) FROM jobs").fetchone() == (1,)
            finally:
                database_check.close()
            print("PASS: isolated Windows Ctrl+C event drains the service and preserves the synthetic database")
        finally:
            if process.poll() is None:
                process.terminate()
                try:
                    process.communicate(timeout=5)
                except subprocess.TimeoutExpired:
                    process.kill()
                    process.communicate()


def check_backup_snapshot_during_uncommitted_write() -> None:
    """A live snapshot must not include part of an uncommitted transaction."""
    with tempfile.TemporaryDirectory(prefix="karsahire-live-backup-") as directory:
        root = Path(directory)
        source = root / "live.sqlite3"
        connection = sqlite3.connect(source)
        try:
            connection.execute(
                "CREATE TABLE batches (id INTEGER PRIMARY KEY, expected_items INTEGER NOT NULL)"
            )
            connection.execute(
                "CREATE TABLE batch_items (batch_id INTEGER NOT NULL REFERENCES batches(id), ordinal INTEGER NOT NULL, PRIMARY KEY(batch_id, ordinal))"
            )
            connection.commit()
        finally:
            connection.close()

        transaction_open = threading.Event()
        allow_commit = threading.Event()
        writer_errors: list[BaseException] = []

        def hold_partial_transaction() -> None:
            writer = sqlite3.connect(source, timeout=5)
            try:
                writer.execute("PRAGMA foreign_keys = ON")
                writer.execute("BEGIN IMMEDIATE")
                writer.execute("INSERT INTO batches(id, expected_items) VALUES(1, 2)")
                writer.execute("INSERT INTO batch_items(batch_id, ordinal) VALUES(1, 1)")
                transaction_open.set()
                if not allow_commit.wait(timeout=5):
                    raise TimeoutError("test did not release the synthetic write transaction")
                writer.execute("INSERT INTO batch_items(batch_id, ordinal) VALUES(1, 2)")
                writer.commit()
            except BaseException as error:  # surface worker failures in the owning test thread
                writer_errors.append(error)
                writer.rollback()
            finally:
                writer.close()

        writer_thread = threading.Thread(target=hold_partial_transaction, daemon=True)
        writer_thread.start()
        try:
            assert transaction_open.wait(timeout=3), "synthetic writer did not open its transaction"
            before_commit = root / "before-commit.sqlite3"
            result = backup_database.create_backup(source, before_commit)
            assert result["integrity_check"] == "ok"
            snapshot = sqlite3.connect(before_commit)
            try:
                assert snapshot.execute("PRAGMA integrity_check").fetchone() == ("ok",)
                assert snapshot.execute("SELECT COUNT(*) FROM batches").fetchone() == (0,)
                assert snapshot.execute("SELECT COUNT(*) FROM batch_items").fetchone() == (0,)
            finally:
                snapshot.close()
        finally:
            allow_commit.set()
            writer_thread.join(timeout=5)
        assert not writer_thread.is_alive(), "synthetic write transaction did not finish"
        assert not writer_errors, f"synthetic writer failed: {writer_errors!r}"

        after_commit = root / "after-commit.sqlite3"
        result = backup_database.create_backup(source, after_commit)
        assert result["integrity_check"] == "ok"
        snapshot = sqlite3.connect(after_commit)
        try:
            assert snapshot.execute("PRAGMA integrity_check").fetchone() == ("ok",)
            batch = snapshot.execute("SELECT expected_items FROM batches WHERE id=1").fetchone()
            item_count = snapshot.execute(
                "SELECT COUNT(*) FROM batch_items WHERE batch_id=1"
            ).fetchone()[0]
            assert batch == (2,) and item_count == 2
        finally:
            snapshot.close()
    print("PASS: live SQLite backup excludes an uncommitted partial write and captures its full commit")


def check_pipeline_version() -> None:
    with patch.dict(os.environ):
        os.environ.pop(server.PIPELINE_VERSION_ENV, None)
        version = server.pipeline_version()
        assert re.fullmatch(r"local-[a-f0-9]{58}", version)
        assert server.pipeline_version() == version
    with patch.dict(os.environ, {server.PIPELINE_VERSION_ENV: "release-2026.09"}):
        assert server.pipeline_version() == "release-2026.09"
    with patch.dict(os.environ, {server.PIPELINE_VERSION_ENV: "unsafe/version"}):
        assert server.pipeline_version() == "unversioned"
    print("PASS: local pipeline fingerprint is stable and controlled build override is validated")


def check_manual_upload_config() -> None:
    assert not server.manual_uploads_enabled_from_env("")
    assert not server.manual_uploads_enabled_from_env("false")
    assert not server.manual_uploads_enabled_from_env("off")
    for enabled_value in ("true", "1", "yes", "on"):
        assert server.manual_uploads_enabled_from_env(enabled_value)
    try:
        server.manual_uploads_enabled_from_env("maybe")
    except SystemExit as error:
        assert server.MANUAL_UPLOADS_ENV in str(error)
    else:
        raise AssertionError("invalid manual upload configuration should stop startup")
    print("PASS: manual CV uploads default off and invalid configuration fails closed")


def run_smoke() -> None:
    check_pipeline_version()
    check_manual_upload_config()
    check_local_bind_config()
    check_shutdown_drains_active_handler()
    check_database_path_config()
    check_shutdown_grace_config()
    check_schema_migration_and_database_backup()
    check_local_restore_cutover_and_rollback()
    check_process_restore_cutover_and_rollback()
    check_windows_ctrl_c_process_shutdown()
    check_backup_snapshot_during_uncommitted_write()
    with tempfile.TemporaryDirectory(prefix="karsahire-smoke-") as temp_dir:
        temp_path = Path(temp_dir)
        server.DATA = temp_path / "data"
        server.DB_PATH = server.DATA / "smoke.sqlite3"
        server.MANUAL_UPLOADS_ENABLED = False
        server.init_db()

        httpd = server.ThreadingHTTPServer(("127.0.0.1", 0), server.Handler)
        thread = threading.Thread(target=httpd.serve_forever, daemon=True)
        thread.start()
        base = f"http://127.0.0.1:{httpd.server_address[1]}"
        parsed_base = urlsplit(base)
        try:
            check_readiness_during_drain(base)
            check_request_guards(base)
            check_request_id_correlation(base)
            check_incomplete_request_header_timeout(base)
            check_absolute_request_header_deadline(base)
            check_request_header_limits()
            check_request_body_deadline()
            check_same_origin_mutation_guard(base)
            check_local_image_normalization()
            check_text_document_extraction()
            check_document_resource_limits()
            check_pdf_worker_process_bounds()
            check_evidence_redaction()
            check_negated_skill_evidence()
            check_candidate_score_ordering()
            check_experience_duration_parsing()
            check_education_mention_parsing()
            check_request_log_redaction()
            check_scanned_pdf_ocr_gate()
            check_scanned_pdf_ocr_stub()
            check_pdf_ocr_image_budget()
            check_ocr_response_limit()
            check_ocr_redirect_block()
            check_ocr_transport_security()
            check_ocr_origin_allowlist()
            check_ocr_configuration_health(base)
            check_ocr_evaluator_no_network()
            check_ocr_preflight_no_network()
            check_windows_local_ocr_wrapper()
            check_synthetic_ocr_corpus_aggregation()
            check_synthetic_format_roundtrip()
            check_release_preflight_environment_filter()
            check_tessdata_fingerprinting()
            check_synthetic_baseline_breakdown()
            check_chart_color_contrast()
            check_reduced_motion_contract()
            check_reviewer_form_accessibility()
            check_annotation_evaluator()
            check_blind_annotation_packets()
            check_blind_ranking_packets()
            status, job = request_json(
                f"{base}/api/jobs",
                "POST",
                {
                    "title": "Smoke Test Engineer",
                    "department": "Local QA",
                    "description": "Build Python services.",
                    "criteria": [{"type": "required", "label": "Python"}],
                },
            )
            assert status == 201, f"job creation returned HTTP {status}"
            job_id = job["id"]

            status, health = request_json(f"{base}/api/health")
            assert status == 200 and health["manual_uploads_enabled"] is False

            with patch.object(server, "multipart_file", side_effect=AssertionError("disabled uploads must not parse the request body")):
                try:
                    upload_text(
                        f"{base}/api/jobs/{job_id}/candidates",
                        "synthetic.txt",
                        "Python engineer with 5 years of experience building Python services.",
                    )
                except urllib.error.HTTPError as error:
                    body = json.loads(error.read().decode("utf-8"))
                    assert error.code == 403 and "dinonaktifkan" in body["error"]
                else:
                    raise AssertionError("manual upload should be disabled by default")

            server.MANUAL_UPLOADS_ENABLED = True
            try:
                upload_text(
                    f"{base}/api/jobs/{job_id}/candidates",
                    "synthetic.txt",
                    "Python engineer with 5 years of experience building Python services.",
                )
            except urllib.error.HTTPError as error:
                error.read()
                assert error.code == 409, f"unapproved upload returned HTTP {error.code}"
            else:
                raise AssertionError("upload should be blocked before both approvals")

            for reviewer, role in (("Recruiter Smoke", "recruiter"), ("Manager Smoke", "hiring_manager")):
                status, _ = request_json(
                    f"{base}/api/jobs/{job_id}/approvals",
                    "POST",
                    {"reviewer": reviewer, "role": role},
                )
                assert status == 201, f"{role} approval returned HTTP {status}"

            status, approved_job = request_json(f"{base}/api/jobs/{job_id}")
            assert status == 200 and approved_job["status"] == "criteria_approved"

            server.MANUAL_UPLOADS_ENABLED = False
            with patch.object(server, "multipart_file", side_effect=AssertionError("disabled uploads must not parse the request body")):
                try:
                    upload_text(
                        f"{base}/api/jobs/{job_id}/candidates",
                        "synthetic-disabled.txt",
                        "Python engineer with 5 years of experience building Python services.",
                    )
                except urllib.error.HTTPError as error:
                    body = json.loads(error.read().decode("utf-8"))
                    assert error.code == 403 and "dinonaktifkan" in body["error"]
                else:
                    raise AssertionError("approved jobs must still reject manual files while the feature is disabled")
            status, unchanged_job = request_json(f"{base}/api/jobs/{job_id}")
            assert status == 200 and not unchanged_job["candidates"], "disabled upload must not create a candidate"
            released_slots = 0
            try:
                for _ in range(server.MAX_CONCURRENT_UPLOADS):
                    assert server.UPLOAD_SEMAPHORE.acquire(blocking=False), "disabled upload must not consume a parser slot"
                    released_slots += 1
            finally:
                for _ in range(released_slots):
                    server.UPLOAD_SEMAPHORE.release()
            server.MANUAL_UPLOADS_ENABLED = True

            files_before_invalid_upload = {
                path.relative_to(server.DATA).as_posix()
                for path in server.DATA.rglob("*")
                if path.is_file()
            }
            try:
                upload_bytes(
                    f"{base}/api/jobs/{job_id}/candidates",
                    "synthetic-broken.pdf",
                    b"%PDF-1.4\n%%EOF\n",
                )
            except urllib.error.HTTPError as error:
                body = json.loads(error.read().decode("utf-8"))
                assert error.code == 400 and body["error"] == "PDF rusak atau tidak dapat dibaca."
            else:
                raise AssertionError("malformed PDF upload should return a client validation error")
            status, unchanged_job = request_json(f"{base}/api/jobs/{job_id}")
            assert status == 200 and not unchanged_job["candidates"], (
                "malformed file must not create a candidate"
            )
            files_after_invalid_upload = {
                path.relative_to(server.DATA).as_posix()
                for path in server.DATA.rglob("*")
                if path.is_file()
            }
            assert files_after_invalid_upload == files_before_invalid_upload, (
                "malformed file must not be persisted under the application data directory"
            )
            released_slots = 0
            try:
                for _ in range(server.MAX_CONCURRENT_UPLOADS):
                    assert server.UPLOAD_SEMAPHORE.acquire(blocking=False), "failed upload should release its slot"
                    released_slots += 1
            finally:
                for _ in range(released_slots):
                    server.UPLOAD_SEMAPHORE.release()

            held_upload_slots = 0
            try:
                for _ in range(server.MAX_CONCURRENT_UPLOADS):
                    assert server.UPLOAD_SEMAPHORE.acquire(blocking=False)
                    held_upload_slots += 1
                try:
                    upload_text(
                        f"{base}/api/jobs/{job_id}/candidates",
                        "busy-synthetic.txt",
                        "Python engineer.",
                    )
                except urllib.error.HTTPError as error:
                    error.read()
                    assert error.code == 503, f"busy upload returned HTTP {error.code}"
                    assert error.headers.get("Retry-After") == "2", "busy upload should advertise a retry delay"
                else:
                    raise AssertionError("uploads should receive retryable backpressure when all slots are occupied")
            finally:
                for _ in range(held_upload_slots):
                    server.UPLOAD_SEMAPHORE.release()

            status, metrics = request_json(f"{base}/api/metrics")
            assert status == 200 and metrics["scope"] == "process"
            assert metrics["uploads_active"] == 0
            assert metrics["upload_busy_rejections_total"] >= 1
            assert metrics["http_response_bytes_total"] > 0
            assert any(
                row["route"] == "/api/jobs/:id/candidates" and row["status"] == 503
                for row in metrics["http_requests_total"]
            )
            assert not re.search(r"\bcand_[a-f0-9]{12}\b", json.dumps(metrics))

            sensitive_filename = "applicant+jane.doe@example.test.txt"
            upload_log = io.StringIO()
            with redirect_stdout(upload_log):
                status, candidate = upload_text(
                    f"{base}/api/jobs/{job_id}/candidates",
                    sensitive_filename,
                    "Python engineer with 5 years of experience building Python services.",
                )
            assert status == 201, f"approved upload returned HTTP {status}"
            released_slots = 0
            try:
                for _ in range(server.MAX_CONCURRENT_UPLOADS):
                    assert server.UPLOAD_SEMAPHORE.acquire(blocking=False), "successful upload should release its slot"
                    released_slots += 1
            finally:
                for _ in range(released_slots):
                    server.UPLOAD_SEMAPHORE.release()
            candidate_id = candidate["id"]
            assert candidate["score"] == 100.0
            assert sensitive_filename not in json.dumps(candidate)
            assert sensitive_filename not in upload_log.getvalue()
            with server.connect() as db:
                for table in ("jobs", "approvals", "candidates", "evidence", "reviews", "audit_events"):
                    for row in db.execute(f"SELECT * FROM {table}").fetchall():
                        assert sensitive_filename not in repr(tuple(row)), f"uploaded filename was persisted in {table}"

            status, _ = request_json(
                f"{base}/api/candidates/{candidate_id}/reviews",
                "POST",
                {
                    "reviewer": "Recruiter Smoke",
                    "role": "recruiter",
                    "decision": "advance",
                    "note": "Evidence reviewed. Contact qa@example.test phone +1 555 010 2020 DOB: 1994-03-12",
                },
            )
            assert status == 201, f"human review returned HTTP {status}"

            status, reviewed_job = request_json(f"{base}/api/jobs/{job_id}")
            assert status == 200
            saved_note = reviewed_job["candidates"][0]["reviews"][0]["note"]
            assert "qa@example.test" not in saved_note
            assert "555 010 2020" not in saved_note and "1994-03-12" not in saved_note
            assert "[email removed]" in saved_note and "[phone removed]" in saved_note

            delete_response_finished = threading.Event()
            delete_statuses = []
            send_json = server.Handler.send_json

            def observe_delete_rejection(handler, status, payload, extra_headers=None):
                if status == 400:
                    delete_statuses.append((status, handler.close_connection))
                    try:
                        return send_json(handler, status, payload, extra_headers)
                    finally:
                        delete_response_finished.set()
                return send_json(handler, status, payload, extra_headers)

            delete_connection = http.client.HTTPConnection(parsed_base.hostname, parsed_base.port, timeout=5)
            try:
                delete_connection.putrequest("DELETE", f"/api/candidates/{candidate_id}", skip_host=True)
                delete_connection.putheader("Host", f"{parsed_base.hostname}:{parsed_base.port}")
                delete_connection.putheader("Origin", base)
                delete_connection.putheader("X-KarsaHire-Request", "same-origin-ui")
                delete_connection.putheader("Content-Length", "1")
                with patch.object(server.Handler, "send_json", observe_delete_rejection):
                    delete_connection.endheaders()
                    assert delete_response_finished.wait(timeout=3), "DELETE with a body should be rejected before changing candidate state"
            finally:
                delete_connection.close()
            assert delete_statuses == [(400, True)]
            with server.connect() as db:
                assert db.execute("SELECT COUNT(*) FROM candidates WHERE id=?", (candidate_id,)).fetchone()[0] == 1

            status, deletion = request_json(f"{base}/api/candidates/{candidate_id}", "DELETE")
            assert status == 200 and deletion["ok"]
            status, updated_job = request_json(f"{base}/api/jobs/{job_id}")
            assert status == 200 and not updated_job["candidates"], "candidate data should be deleted"
            status, events = request_json(f"{base}/api/jobs/{job_id}/events")
            deletion_events = [event for event in events if event["event_type"] == "candidate_deleted"]
            assert status == 200 and deletion_events
            assert deletion_events[0]["candidate_id"] is None
            assert deletion_events[0]["actor"] == "user"
            assert not any(event["candidate_id"] == candidate_id for event in events)
            with server.connect() as db:
                assert db.execute(
                    "SELECT COUNT(*) FROM audit_events WHERE candidate_id=?", (candidate_id,)
                ).fetchone()[0] == 0, "candidate-linked audit events should be removed with the candidate"

            status, demo_fixture = request_json(
                f"{base}/api/jobs/{job_id}/load-demo-candidate", "POST", {},
            )
            assert status == 201 and demo_fixture["added"] and demo_fixture["id"]
            status, repeated_demo = request_json(
                f"{base}/api/jobs/{job_id}/load-demo-candidate", "POST", {},
            )
            assert status == 200 and not repeated_demo["added"]
            status, demo_job = request_json(f"{base}/api/jobs/{job_id}")
            assert status == 200 and len(demo_job["candidates"]) == 1
            with server.connect() as db:
                demo_profile = json.loads(db.execute(
                    "SELECT profile_json FROM candidates WHERE id=?", (demo_fixture["id"],)
                ).fetchone()[0])
            assert demo_profile["dataset_source"] == server.BUILTIN_DEMO_SOURCE_KEY

            status, corpus_job = request_json(
                f"{base}/api/jobs",
                "POST",
                {
                    "title": "Synthetic Corpus Smoke",
                    "department": "Local QA",
                    "description": "Validate the local synthetic parser corpus import.",
                    "criteria": [{"type": "required", "label": "Python"}],
                },
            )
            assert status == 201, f"corpus requisition creation returned HTTP {status}"
            corpus_job_id = corpus_job["id"]
            for reviewer, role in (("Corpus Recruiter", "recruiter"), ("Corpus Manager", "hiring_manager")):
                status, _ = request_json(
                    f"{base}/api/jobs/{corpus_job_id}/approvals",
                    "POST",
                    {"reviewer": reviewer, "role": role},
                )
                assert status == 201, f"corpus {role} approval returned HTTP {status}"

            status, imported = request_json(f"{base}/api/jobs/{corpus_job_id}/load-synthetic-data", "POST", {})
            assert status == 200 and imported["added"] == server.SYNTHETIC_DATASET_SIZE
            assert imported["skipped"] == 0 and imported["total_candidates"] == server.SYNTHETIC_DATASET_SIZE
            status, repeated = request_json(f"{base}/api/jobs/{corpus_job_id}/load-synthetic-data", "POST", {})
            assert status == 200 and repeated["added"] == 0
            assert repeated["skipped"] == server.SYNTHETIC_DATASET_SIZE
            assert repeated["total_candidates"] == server.SYNTHETIC_DATASET_SIZE

            print("PASS: manual uploads fail closed, then the approval gate and explicit local test opt-in work")
            print("PASS: fixed synthetic demo fixture loads locally and is idempotent")
            print("PASS: unauthenticated startup rejects non-loopback bind and invalid ports")
            print("PASS: SQLite v1 migration, synthetic backup/restore integrity, no-overwrite, future-schema rejection, and repository-destination guard")
            print("PASS: upload parsing has a process-wide concurrency cap and returns retryable HTTP 503 when busy")
            print("PASS: process-local metrics expose bounded route counters and upload backpressure")
            print("PASS: approved text upload, evidence score, and human review")
            print("PASS: PII-shaped source filenames are not returned or persisted")
            print("PASS: candidate deletion removes candidate data and preserves deletion audit")
            print("PASS: browser security headers and JSON request validation")
            print("PASS: liveness stays available while readiness reports an unsupported database schema")
            print("PASS: malformed and ambiguous Content-Length/Transfer-Encoding headers are rejected")
            print("PASS: duplicate Host, Origin, and UI marker headers are rejected")
            print("PASS: request-body total deadline releases the socket and closes slow clients")
            print("PASS: loopback Host guard rejects hostile reads and same-origin guard rejects hostile writes")
            print("PASS: rejected mutation closes without draining an incomplete request body")
            print("PASS: synthetic corpus import and idempotent re-import (32 records)")
            print("PASS: PNG/JPEG upload paths, normalization, and filename/content validation with stubbed OCR")
            print("PASS: synthetic TXT, text-PDF, and DOCX extraction paths plus safe malformed-document errors")
            print("PASS: malformed synthetic PDF returns HTTP 400 without persisting a file or candidate and releases its slot")
            print("PASS: image-only PDF is gated on OCR configuration without a network call")
            print("PASS: OCR response size is bounded using a stubbed response (no network)")
            print("PASS: OCR redirects are rejected before credentials can follow them")
            print("PASS: OCR requires HTTPS, a clean URL, and an exact approved origin (stubbed transport)")
            print("PASS: OCR origin allowlist fails closed and normalizes equivalent HTTPS origins")
            print("PASS: health reports complete OCR configuration without exposing the key or URL")
            print("PASS: PDF page/image-count, raster-image, DOCX expansion, and empty-text bounds")
            print("PASS: contact details and date of birth are removed from evidence and review notes")
            print("PASS: experience duration recognizes synthetic Indonesian year phrases and skips explicitly labeled ages")
            print("PASS: synthetic matched/partial/unknown and negated evidence snippets keep status and page attribution")
            print("PASS: request logs redact candidate IDs and query parameters")
            print("PASS: structured request logs include response bytes without retaining client IP")
            print("PASS: OCR WER opt-in path using a stubbed synthetic transcription (no network)")
            print("PASS: OCR evaluation preflight makes no extraction or network call without opt-in")
            print("PASS: synthetic baseline unmatched-label breakdown matches aggregate count")
            print("PASS: chart text/fill contrast and non-color status labels remain available")
            print("PASS: annotation, citation-reference agreement, criterion metrics, and candidate-participation privacy gates")
        finally:
            httpd.shutdown()
            httpd.server_close()
            thread.join(timeout=2)


if __name__ == "__main__":
    try:
        run_smoke()
    except (AssertionError, OSError, ValueError) as error:
        print(f"Smoke test failed: {error}", file=sys.stderr)
        raise SystemExit(1)
